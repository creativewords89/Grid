"""Read Excel and CSV files into tables (SPEC section 6.1).

- Excel values are the results Excel saved for formulas (`data_only=True`). A file whose
  formulas were never calculated (e.g. made by a script) gets a warning.
- The header row is the first row that looks like column names: at least two filled
  cells (or the only column), all text, and at least half as wide as the widest of the
  first rows. Title lines above it are kept as notes. Without such a row, the first
  non-empty row is the header.
- Empty rows and columns are skipped. Hidden sheets, rows and columns are included.
"""

import csv
import io
import re
from collections.abc import Iterable, Iterator, Sequence
from contextlib import contextmanager
from datetime import date, datetime, time
from itertools import zip_longest
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.files.sheets import Sheet, Workbook
from app.jobs.queue import PermanentError

DAMAGED_XLSX = "This Excel file is damaged and can't be read. Try saving it again as .xlsx."
DAMAGED_CSV = "This CSV file can't be read. Save it again as CSV (UTF-8) and upload it."
UNCALCULATED = (
    "Some formulas have no saved results, so those cells were skipped. Open the file in "
    "Excel, save it, then upload it again."
)
HEADER_SCAN_ROWS = 10
_NUMBER = re.compile(r"^[-+]?[$£€]?\s?[\d,]*\.?\d+%?$")

Cells = list[str]


# --- turning cell values into text -----------------------------------------------------


_DECIMALS = re.compile(r"\.(0+)")


def _number(value: float, number_format: str) -> str:
    """Roughly as Excel shows it: thousands separators, fixed decimals, %, currency."""
    section = number_format.split(";")[0]  # the format for positive numbers
    if section == "General" or not any(ch in section for ch in "0#"):
        text = f"{value:.10g}"
        if "e" in text and abs(value) < 1e15:  # .10g switches to exponents for big numbers
            text = f"{value:.2f}".rstrip("0").rstrip(".")
        return text
    decimals = len(match.group(1)) if (match := _DECIMALS.search(section)) else 0
    percent = "%" in section
    shown = value * 100 if percent else value
    grouping = "," if "," in section.split(".")[0] else ""
    text = f"{abs(shown):{grouping}.{decimals}f}"
    symbol = next((sym for sym in ("$", "£", "€") if sym in section), "")
    sign = "-" if shown < 0 and float(text.replace(",", "")) != 0 else ""
    return f"{sign}{symbol}{text}{'%' if percent else ''}"


def cell_text(value: Any, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, int | float):
        return _number(value, number_format or "General")
    return " ".join(str(value).split())


# --- finding the header row and naming columns ------------------------------------------


def _filled(row: Cells) -> int:
    return sum(1 for cell in row if cell)


def _looks_like_header(row: Cells, widest: int) -> bool:
    filled = [cell for cell in row if cell]
    if not filled:
        return False
    if any(_NUMBER.match(cell.replace(" ", "")) for cell in filled):
        return False
    return (len(filled) >= 2 or widest == 1) and len(filled) * 2 >= widest


def _column_names(header: Cells, width: int) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        name = header[index] if index < len(header) and header[index] else ""
        name = name or f"Column {get_column_letter(index + 1)}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return names


def build_sheet(name: str | None, numbered_rows: Iterable[tuple[int, Cells]]) -> Sheet | None:
    """Turn a sheet's rows (row number, cell texts) into a Sheet, or None when empty."""
    rows = [(n, cells) for n, cells in numbered_rows if any(cells)]
    if not rows:
        return None
    first = rows[:HEADER_SCAN_ROWS]
    widest = max(_filled(cells) for _, cells in first)
    header_at = next(
        (i for i, (_, cells) in enumerate(first) if _looks_like_header(cells, widest)), 0
    )
    notes = [" ".join(cell for cell in cells if cell) for _, cells in rows[:header_at]]
    header = rows[header_at][1]
    body = rows[header_at + 1 :]

    width = max(len(cells) for _, cells in rows[header_at:])
    used = [i for i in range(width) if any(i < len(c) and c[i] for _, c in rows[header_at:])]
    names = _column_names(header, width)
    return Sheet(
        name=name,
        headers=[names[i] for i in used],
        rows=[(n, [cells[i] if i < len(cells) else "" for i in used]) for n, cells in body],
        notes=notes,
    )


# --- Excel -----------------------------------------------------------------------------


def _excel_rows(sheet: Any) -> Iterator[tuple[int, Cells]]:
    for row in sheet.iter_rows():
        cells = [cell_text(c.value, getattr(c, "number_format", "General")) for c in row]
        number = next((c.row for c in row if getattr(c, "row", None)), None)
        if number is not None:
            yield number, cells


def _uncalculated(path: Path, values: Any) -> bool:
    """True when a formula cell has no saved result."""
    with _open(path, data_only=False) as formulas:
        for name in formulas.sheetnames:
            pairs = zip_longest(
                formulas[name].iter_rows(values_only=True),
                values[name].iter_rows(values_only=True),
                fillvalue=(),
            )
            for formula_row, value_row in pairs:
                for formula, value in zip_longest(formula_row, value_row):
                    if isinstance(formula, str) and formula.startswith("=") and value is None:
                        return True
    return False


@contextmanager
def _open(path: Path, data_only: bool) -> Iterator[Any]:
    # Opened from a file object: openpyxl refuses a path whose name doesn't end in .xlsx,
    # and uploads are stored under random names. Both are closed afterwards.
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=data_only)
        try:
            yield workbook
        finally:
            workbook.close()


def extract_xlsx(path: Path) -> Workbook:
    try:
        with _open(path, data_only=True) as workbook:
            sheets: list[Sheet] = []
            for name in workbook.sheetnames:  # hidden sheets included
                ws = workbook[name]
                if not hasattr(ws, "iter_rows"):  # a chart sheet
                    continue
                sheet = build_sheet(name, _excel_rows(ws))
                if sheet is not None:
                    sheets.append(sheet)
            warning = UNCALCULATED if _uncalculated(path, workbook) else None
    except Exception as exc:  # openpyxl raises many types for a broken file
        raise PermanentError(DAMAGED_XLSX) from exc
    return Workbook(sheets, warning)


# --- CSV -------------------------------------------------------------------------------


def _decode(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")  # Excel's "CSV" on Windows


def _dialect(sample: str) -> type[csv.Dialect] | csv.Dialect:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def extract_csv(path: Path) -> Workbook:
    text = _decode(path.read_bytes())
    try:
        reader = csv.reader(io.StringIO(text, newline=""), _dialect(text[: 64 * 1024]))
        rows: Sequence[tuple[int, Cells]] = [
            (number, [cell_text(cell) for cell in row]) for number, row in enumerate(reader, 1)
        ]
    except csv.Error as exc:
        raise PermanentError(DAMAGED_CSV) from exc
    sheet = build_sheet(None, rows)
    return Workbook([sheet] if sheet else [])
