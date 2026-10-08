"""Build real PDF and Word files for tests."""

from collections.abc import Callable
from pathlib import Path

import docx
import pymupdf
from docx.enum.text import WD_BREAK

LOREM = (
    "Our local SEO plans include a monthly report for every client. Each report shows "
    "rankings, calls and direction requests. Reviews are answered within two working days."
)

PageWriter = Callable[[pymupdf.Page], object]


def pdf(path: Path, *pages: PageWriter, **save: object) -> Path:
    doc = pymupdf.open()
    for write in pages:
        write(doc.new_page())
    doc.save(path, **save)
    doc.close()
    return path


def text(page: pymupdf.Page, y: float, value: str, size: float = 11, bold: bool = False) -> float:
    """Write a block of text from height y; returns the y below it."""
    rect = pymupdf.Rect(72, y, 520, y + 400)
    left = page.insert_textbox(rect, value, fontsize=size, fontname="hebo" if bold else "helv")
    return rect.y1 - left + 8


def table(page: pymupdf.Page, y: float, rows: list[list[str]]) -> float:
    """A grid of ruled cells, as most exported PDFs draw tables."""
    width, height = 150, 22
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = pymupdf.Rect(
                72 + c * width, y + r * height, 72 + (c + 1) * width, y + (r + 1) * height
            )
            page.draw_rect(cell, color=(0, 0, 0), width=0.8)
            page.insert_text((cell.x0 + 4, cell.y1 - 7), value, fontsize=10)
    return y + len(rows) * height + 12


def locked_pdf(path: Path) -> Path:
    """A PDF that needs a password to open."""
    aes_256 = pymupdf.PDF_ENCRYPT_AES_256  # type: ignore[attr-defined]
    return pdf(
        path,
        lambda p: text(p, 60, LOREM),
        encryption=aes_256,
        user_pw="secret",
        owner_pw="owner",
    )


def image_page(page: pymupdf.Page) -> None:
    """A full-page picture with no text: what a scanner produces."""
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), False)
    pixmap.set_rect(pixmap.irect, (200, 200, 200))
    page.insert_image(page.rect, pixmap=pixmap)


def word(
    path: Path, build: Callable[[docx.document.Document], object], pages: int | None = None
) -> Path:
    document = docx.Document()
    build(document)
    document.save(str(path))
    if pages is not None:
        _set_page_count(path, pages)
    return path


def page_break(document: docx.document.Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _set_page_count(path: Path, pages: int) -> None:
    """Word records the page count in docProps/app.xml when it saves a file."""
    import zipfile

    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = data.replace(b"<Pages>1</Pages>", f"<Pages>{pages}</Pages>".encode())
                if b"<Pages>" not in data:
                    data = data.replace(
                        b"</Properties>", f"<Pages>{pages}</Pages></Properties>".encode()
                    )
            dst.writestr(item, data)
    tmp.replace(path)
