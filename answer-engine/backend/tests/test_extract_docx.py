"""Reading Word documents into blocks (SPEC section 6.1)."""

from pathlib import Path

import docx.document
import pytest

from app.files.extract_docx import DAMAGED, extract_docx
from app.jobs.queue import PermanentError
from tests.files.build import page_break, word


def test_headings_paragraphs_lists_and_tables(tmp_path: Path) -> None:
    def build(d: docx.document.Document) -> None:
        d.add_heading("Onboarding", level=0)
        d.add_paragraph("Welcome to GridRankers.")
        d.add_heading("First week", level=1)
        d.add_paragraph("Set up the Google Business Profile.", style="List Bullet")
        d.add_paragraph("Collect logins.", style="List Number")
        d.add_heading("Who does what", level=2)
        t = d.add_table(rows=3, cols=2)
        for r, (a, b) in enumerate(
            [("Task", "Owner"), ("Audit", "Sara | Ali"), ("Report", "Nadia")]
        ):
            t.cell(r, 0).text, t.cell(r, 1).text = a, b
        d.add_paragraph("")  # empty paragraphs are skipped

    result = extract_docx(word(tmp_path / "a.docx", build))

    assert [(b.kind, b.level, b.text) for b in result.blocks if b.kind != "table"] == [
        ("heading", 0, "Onboarding"),
        ("text", 1, "Welcome to GridRankers."),
        ("heading", 1, "First week"),
        ("text", 1, "- Set up the Google Business Profile."),
        ("text", 1, "- Collect logins."),
        ("heading", 2, "Who does what"),
    ]
    [table] = [b for b in result.blocks if b.kind == "table"]
    assert table.text.splitlines() == [
        "| Task | Owner |",
        "|---|---|",
        "| Audit | Sara \\| Ali |",
        "| Report | Nadia |",
    ]


def test_merged_cells_appear_once(tmp_path: Path) -> None:
    def build(d: docx.document.Document) -> None:
        t = d.add_table(rows=2, cols=3)
        merged = t.cell(0, 0).merge(t.cell(0, 2))
        merged.text = "Fees 2026"
        for c, value in enumerate(["Basic", "Pro", "Agency"]):
            t.cell(1, c).text = value

    [table] = extract_docx(word(tmp_path / "a.docx", build)).blocks

    assert table.text.splitlines() == [
        "| Fees 2026 |  |  |",
        "|---|---|---|",
        "| Basic | Pro | Agency |",
    ]


def test_page_numbers_follow_page_breaks(tmp_path: Path) -> None:
    def build(d: docx.document.Document) -> None:
        d.add_paragraph("On page one.")
        page_break(d)
        d.add_paragraph("On page two.")
        page_break(d)
        d.add_paragraph("On page three.")

    result = extract_docx(word(tmp_path / "a.docx", build, pages=3))

    assert [(b.page, b.text) for b in result.blocks] == [
        (1, "On page one."),
        (2, "On page two."),
        (3, "On page three."),
    ]
    assert result.page_count == 3


def test_without_page_information_pages_are_unknown(tmp_path: Path) -> None:
    result = extract_docx(word(tmp_path / "a.docx", lambda d: d.add_paragraph("Hello.")))

    assert [b.page for b in result.blocks] == [None]


def test_damaged_word_file_fails_for_good(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04 not a real word file")

    with pytest.raises(PermanentError) as refused:
        extract_docx(path)
    assert str(refused.value) == DAMAGED
