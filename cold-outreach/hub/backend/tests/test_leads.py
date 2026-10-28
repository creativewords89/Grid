"""The Leads screen and what agents save for leads (SPEC sections 6.2 and 6.4)."""

import io
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import AuditEntry, DoNotContact, Lead, Role, User
from tests.conftest import UserFactory, sign_in
from tests.hub_data import batch_id, put_bot
from tests.test_agents import MANIFEST, add_agent, agent_call
from tests.test_imports import FILE, STANDARD, row, upload, workbook


@pytest.fixture
def owner(client: TestClient, make_user: UserFactory) -> User:
    user = make_user(Role.OWNER, name="Olivia")
    sign_in(client, user.email)
    return user


@pytest.fixture
def reports(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(get_settings(), "reports_dir", str(tmp_path))
    return tmp_path


@pytest.fixture
def saved(client: TestClient, owner: User) -> None:
    """FILE in Batch 1 (L-000001 … L-000004, L-000002 marked Top) and one more lead in
    Batch 2 (L-000005)."""
    for content, batch in (
        (FILE, "Batch 1"),
        (
            {"Austin": [STANDARD, row("Shine ATX", "https://share.google/atx", "512 555 0100")]},
            "Batch 2",
        ),
    ):
        preview = upload(client, workbook(content), batch=batch)
        assert client.post(f"/api/imports/{preview['id']}/confirm").status_code == 200
    # Top is chosen in the Hub (Top Shine).
    marked = client.post("/api/leads/top", json={"lead_ids": ["L-000002"], "top": True})
    assert marked.json() == {"changed": 1}


def page(client: TestClient, **params: Any) -> dict[str, Any]:
    response = client.get("/api/leads", params=params)
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()
    return body


def ids(body: dict[str, Any]) -> list[str]:
    return [lead["lead_id"] for lead in body["leads"]]


# --- looking ------------------------------------------------------------------------------


def test_filters_and_search(client: TestClient, saved: None) -> None:
    every = page(client)
    assert (every["total"], every["pages"]) == (5, 1)
    assert ids(page(client, batch_id=batch_id(client, "Batch 2"))) == ["L-000005"]
    assert ids(page(client, q="multi")) == ["L-000001"]
    assert ids(page(client, q="918-523-8300")) == ["L-000001"]
    assert ids(page(client, q="kevin@")) == ["L-000001"]
    assert ids(page(client, city="dallas")) == ["L-000004"]
    assert ids(page(client, level="top")) == ["L-000002"]
    assert "L-000002" not in ids(page(client, level="not_top"))
    assert ids(page(client, sort="newest"))[0] == "L-000005"
    no_site = page(client, website="no")
    assert all(not lead["data"].get("website") for lead in no_site["leads"])


def test_pages_of_one_hundred(
    client: TestClient, saved: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.api import leads

    monkeypatch.setattr(leads, "PAGE", 2)

    second = page(client, page=2)

    assert (second["total"], second["pages"], ids(second)) == (5, 3, ["L-000003", "L-000004"])


def test_one_lead_with_its_files(client: TestClient, saved: None) -> None:
    lead = client.get("/api/leads/L-000001").json()

    assert (lead["business_name"], lead["batch_name"], lead["files"]) == (
        "Multi-Clean",
        "Batch 1",
        [],
    )
    assert lead["can_edit"] is True
    assert client.get("/api/leads/L-999999").status_code == 404


# --- the Excel file -----------------------------------------------------------------------


def test_the_excel_file_has_the_filtered_leads_and_text_stays_text(
    client: TestClient, saved: None
) -> None:
    response = client.get("/api/leads/export", params={"batch_id": batch_id(client)})

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    assert header[:6] == ("Lead ID", "Batch", "Niche", "City", "Level", "Business Name")
    assert [r[0] for r in rows[1:]] == ["L-000001", "L-000002", "L-000003", "L-000004"]
    # A business name like "=1+1" or a cell like =HYPERLINK(...) must never run in Excel.
    formulas = [c for r in sheet.iter_rows() for c in r if c.data_type == "f"]
    assert formulas == []


def test_a_formula_like_name_stays_text(client: TestClient, owner: User) -> None:
    content = workbook({"Tulsa": [STANDARD, row("=1+1", "https://share.google/f", "1")]})
    preview = upload(client, content)
    client.post(f"/api/imports/{preview['id']}/confirm")

    sheet = load_workbook(io.BytesIO(client.get("/api/leads/export").content)).active
    assert sheet is not None
    cell = sheet["F2"]
    assert (cell.value, cell.data_type) == ("=1+1", "s")


# --- changing -----------------------------------------------------------------------------


def test_staff_change_a_lead_and_top_protects_it(
    client: TestClient, saved: None, db: Session
) -> None:
    changed = client.patch(
        "/api/leads/L-000001",
        json={
            "email": "Owner@MC.com, kevin@mc.com",
            "phone": "(918) 555-7777",
            "website": "https://mc.com/?utm_source=x",
            "level": "Top",
            "notes": "Called once",
        },
    ).json()

    assert changed["emails"] == ["owner@mc.com", "kevin@mc.com"]
    assert changed["data"]["phone"] == "(918) 555-7777"
    assert changed["data"]["website"] == "https://mc.com/"
    assert (changed["level"], changed["data"]["notes"]) == ("Top", "Called once")
    blocked = {
        d.value for d in db.scalars(select(DoNotContact).where(DoNotContact.lead_id == "L-000001"))
    }
    assert blocked == {"owner@mc.com", "kevin@mc.com", "9185557777", "https://share.google/mc"}

    client.patch("/api/leads/L-000001", json={"level": ""})
    db.expire_all()
    assert db.scalars(select(DoNotContact).where(DoNotContact.lead_id == "L-000001")).all() == []
    history = client.get("/api/history", params={"entity": "lead"}).json()["entries"]
    edits = [e for e in history if e["action"] == "edit"]
    assert edits[-1]["changes"][0]["field"] == "email"


def test_a_bad_email_is_refused(client: TestClient, saved: None) -> None:
    response = client.patch("/api/leads/L-000001", json={"email": "kevin@"})

    assert (response.status_code, response.json()["error"]["fields"]) == (
        422,
        {"email": "This doesn't look like an email: kevin@"},
    )


def test_delete_a_lead_but_keep_it_off_limits(client: TestClient, saved: None, db: Session) -> None:
    assert client.delete("/api/leads/L-000002").status_code == 204

    assert client.get("/api/leads/L-000002").status_code == 404
    left = db.scalars(select(DoNotContact.value).where(DoNotContact.lead_id == "L-000002")).all()
    assert "boss@ts.com" in left


def test_viewers_look_and_download_but_cannot_change(
    client: TestClient, saved: None, make_user: UserFactory
) -> None:
    client.cookies.clear()
    sign_in(client, make_user(Role.VIEWER).email)

    assert page(client)["total"] == 5
    assert client.get("/api/leads/export").status_code == 200
    assert client.get("/api/leads/L-000001").json()["can_edit"] is False
    assert client.patch("/api/leads/L-000001", json={"notes": "x"}).status_code == 403
    assert client.delete("/api/leads/L-000001").status_code == 403


def test_signed_out_people_see_nothing(client: TestClient) -> None:
    assert client.get("/api/leads").status_code == 401
    assert client.get("/api/leads/export").status_code == 401


# --- what agents save ---------------------------------------------------------------------


def test_an_agent_reads_full_rows_and_saves_its_result(client: TestClient, saved: None) -> None:
    token = add_agent(client)
    agent_call(client, "PUT", "/manifest", token, MANIFEST)
    batch = batch_id(client)
    put_bot(client, batch, "gbp-audit")
    client.put(f"/api/batches/{batch}/cities", json={"active": True})  # small cities

    rows = agent_call(client, "GET", f"/batches/{batch}/leads", token).json()
    multi = rows[0]
    assert (multi["business_name"], multi["phone"], multi["result"]) == (
        "Multi-Clean",
        "(918) 523-8300",
        None,
    )
    saved_result = agent_call(
        client,
        "PUT",
        "/leads/L-000001/result",
        token,
        {
            "status": "OK",
            "score": 61,
            "summary": "Posts are rare.",
            "details": {"posts_90_days": 1},
            "benchmark_version": 2,
        },
    )

    assert saved_result.status_code == 204
    again = agent_call(client, "GET", f"/batches/{batch}/leads", token).json()[0]
    assert (again["result"]["score"], again["result"]["benchmark_version"]) == (61, 2)
    shown = page(client, sort="score", score_of="gbp-audit")
    assert shown["leads"][0]["results"]["gbp-audit"]["summary"] == "Posts are rare."
    assert shown["agents"] == {"gbp-audit": "GBP"}
    sheet = load_workbook(io.BytesIO(client.get("/api/leads/export").content)).active
    assert sheet is not None
    header = next(sheet.iter_rows(values_only=True))
    assert header[-4:] == ("GBP Status", "GBP Score", "GBP Summary", "GBP Checked")


def test_bad_results_are_refused(client: TestClient, saved: None) -> None:
    token = add_agent(client)

    assert (
        agent_call(client, "PUT", "/leads/L-000001/result", token, {"status": ""}).status_code
        == 422
    )
    assert (
        agent_call(client, "PUT", "/leads/L-000001/result", token, {"status": "OK", "score": 150})
    ).status_code == 422
    assert (
        agent_call(client, "PUT", "/leads/L-999999/result", token, {"status": "OK"}).status_code
        == 404
    )


def save_file(client: TestClient, token: str, name: str, body: bytes, kind: str) -> Any:
    return client.post(
        "/api/agent/v1/leads/L-000001/files",
        files={"file": (name, body, kind)},
        data={"name": name},
        headers={"Authorization": f"Bearer {token}"},
    )


def test_an_agent_saves_a_report_and_people_download_it(
    client: TestClient, saved: None, reports: Path
) -> None:
    token = add_agent(client)

    first = save_file(client, token, "gbp-report.pdf", b"%PDF-1 first", "application/pdf")
    second = save_file(client, token, "gbp-report.pdf", b"%PDF-1 second", "application/pdf")

    assert first.status_code == 201 and first.json() == second.json()  # replaced, same id
    [listed] = client.get("/api/leads/L-000001").json()["files"]
    assert (listed["name"], listed["agent_key"], listed["size"]) == (
        "gbp-report.pdf",
        "gbp-audit",
        13,
    )
    download = client.get(f"/api/leads/L-000001/files/{listed['id']}")
    assert download.content == b"%PDF-1 second"
    assert download.headers["content-disposition"].startswith("attachment")
    assert (reports / "L-000001" / listed["id"]).exists()

    assert client.delete("/api/leads/L-000001").status_code == 204
    assert not (reports / "L-000001" / listed["id"]).exists()


@pytest.mark.parametrize(
    ("name", "kind", "code"),
    [
        ("page.html", "text/html", 415),
        ("../../etc/passwd", "application/pdf", 422),
        (".hidden.pdf", "application/pdf", 422),
    ],
)
def test_unsafe_files_are_refused(
    client: TestClient, saved: None, reports: Path, name: str, kind: str, code: int
) -> None:
    token = add_agent(client)

    assert save_file(client, token, name, b"x", kind).status_code == code
    assert list(reports.iterdir()) == []


def test_another_leads_file_id_is_not_found(client: TestClient, saved: None, reports: Path) -> None:
    token = add_agent(client)
    file_id = save_file(client, token, "r.json", b"{}", "application/json").json()["id"]

    assert client.get(f"/api/leads/L-000002/files/{file_id}").status_code == 404


# --- choosing Top (step 17) -------------------------------------------------------------------


def test_many_leads_are_marked_top_and_back(client: TestClient, saved: None, db: Session) -> None:
    marked = client.post(
        "/api/leads/top", json={"lead_ids": ["L-000001", "L-000003", "L-000002"], "top": True}
    )

    assert marked.json() == {"changed": 2}  # L-000002 already was
    assert ids(page(client, level="top")) == ["L-000001", "L-000002", "L-000003"]
    blocked = set(db.scalars(select(DoNotContact.lead_id)))
    assert blocked >= {"L-000001", "L-000002"}
    entry = db.scalar(
        select(AuditEntry).where(AuditEntry.action == "mark_top").order_by(AuditEntry.at.desc())
    )
    assert entry is not None and entry.title == "2 leads"

    client.post("/api/leads/top", json={"lead_ids": ["L-000001"], "top": False})
    assert "L-000001" not in ids(page(client, level="top"))
    assert "L-000001" not in set(db.scalars(select(DoNotContact.lead_id)))


def test_top_needs_staff_and_at_most_500(
    client: TestClient, saved: None, make_user: UserFactory
) -> None:
    too_many = client.post("/api/leads/top", json={"lead_ids": ["L-1"] * 501, "top": True})
    assert too_many.status_code == 422
    client.cookies.clear()
    sign_in(client, make_user(Role.VIEWER).email)

    refused = client.post("/api/leads/top", json={"lead_ids": ["L-000001"], "top": True})
    assert refused.status_code == 403


def test_sort_by_rating_and_reviews_with_odd_values(
    client: TestClient, saved: None, db: Session
) -> None:
    for lead_id, rating, reviews in (
        ("L-000001", "4.9", "1,204"),
        ("L-000003", "n/a", "88"),
        ("L-000004", "3.5", ""),
    ):
        lead = db.get(Lead, lead_id)
        assert lead is not None
        lead.data = dict(lead.data or {}) | {"rating": rating, "reviews": reviews}
    db.commit()

    by_rating = ids(page(client, batch_id=batch_id(client), sort="rating"))
    by_reviews = ids(page(client, batch_id=batch_id(client), sort="reviews"))

    assert by_rating[0] == "L-000001"
    assert by_rating.index("L-000004") < by_rating.index("L-000003")  # "n/a" goes last
    assert by_reviews[0] == "L-000001"  # 1,204 is a number
