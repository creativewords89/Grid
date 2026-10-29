"""Lead Check (SPEC steps 26 and 27): each field holds what its column says, a file with an
error holds all of its leads back, the report and its alert, the filters, editing every
field in the Hub and one-click fixes."""

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Alert, Role, User
from app.imports.check import check_lead, check_tab, fix_for
from tests.conftest import UserFactory, sign_in
from tests.hub_data import batch_id, put_bot
from tests.test_imports import STANDARD, row, upload, workbook
from tests.test_pipeline import bot, ready

CLEAN = {
    "business_name": "Multi-Clean",
    "gbp_link": "https://share.google/mc",
    "address": "5874 S 129th E Ave, Tulsa, OK 74134",
    "city": "Tulsa",
    "website": "https://multiclean1.com/",
    "phone": "(918) 523-8300",
    "whatsapp": "+1 918 523 8300",
    "email": "kevin@multiclean1.com",
    "facebook": "https://facebook.com/multiclean",
    "instagram": "https://instagram.com/multiclean",
    "reviews": "1,204",
    "rating": "4.8",
}


def codes(**changes: str) -> list[str]:
    return [i.code for i in check_lead(CLEAN | changes)]


# --- the rules ----------------------------------------------------------------------------------


def test_a_clean_lead_has_no_problem() -> None:
    assert codes() == []
    assert codes(phone="918-523-8300 ext 4", gbp_link="https://maps.google.com/?cid=1") == []


@pytest.mark.parametrize(
    ("changes", "expected"),
    [
        ({"address": "5874 S 129th E Ave (918) 555-0101"}, "ADDRESS_HAS_PHONE"),
        ({"address": "5874 S 129th E Ave, Tulsa, info@mc.com"}, "ADDRESS_HAS_EMAIL"),
        ({"address": "www.mc.com 5874 S 129th E Ave, Tulsa"}, "ADDRESS_HAS_LINK"),
        ({"website": "https://www.facebook.com/multiclean"}, "WEBSITE_IS_FACEBOOK"),
        ({"website": "instagram.com/mc"}, "WEBSITE_IS_INSTAGRAM"),
        ({"website": "https://maps.google.com/?cid=2"}, "WEBSITE_IS_GOOGLE"),
        ({"website": "call us"}, "WEBSITE_NOT_LINK"),
        ({"phone": "call 918 555"}, "PHONE_HAS_LETTERS"),
        ({"phone": "555-0101"}, "PHONE_WRONG_LENGTH"),
        ({"whatsapp": "yes"}, "WHATSAPP_NOT_PHONE"),
        ({"gbp_link": "https://multiclean1.com"}, "GBP_NOT_MAPS"),
        ({"business_name": "https://multiclean1.com"}, "NAME_IS_LINK"),
        ({"business_name": "(918) 523-8300"}, "NAME_IS_PHONE"),
        ({"city": "5874 S 129th"}, "CITY_HAS_NUMBERS"),
        ({"facebook": "https://instagram.com/mc"}, "FACEBOOK_IS_INSTAGRAM"),
        ({"facebook": "mc cleaning"}, "FACEBOOK_NOT_FACEBOOK"),
        ({"instagram": "https://facebook.com/mc"}, "INSTAGRAM_IS_FACEBOOK"),
        ({"reviews": "lots"}, "REVIEWS_NOT_NUMBER"),
        ({"rating": "great"}, "RATING_NOT_NUMBER"),
        ({"rating": "7.5"}, "RATING_OUT_OF_RANGE"),
        ({"reviews": "4.8", "rating": "120"}, "REVIEWS_RATING_SWAPPED"),
    ],
)
def test_each_error(changes: dict[str, str], expected: str) -> None:
    found = check_lead(CLEAN | changes)

    assert expected in [i.code for i in found]
    assert next(i for i in found if i.code == expected).severity == "error"


def test_missing_or_doubtful_data_only_warns() -> None:
    found = check_lead(
        CLEAN | {"gbp_link": "", "phone": "", "address": "1 Main St, Dallas", "email": "test@x.com"}
    )

    assert {i.code for i in found} == {
        "NO_GBP_LINK",
        "NO_PHONE",
        "ADDRESS_OTHER_CITY",
        "EMAIL_PLACEHOLDER",
    }
    assert all(i.severity == "warning" for i in found)
    tab = check_tab(["business_name", "gbp_link", "phone"], ["phone"], 2)
    assert {(i.code, i.field) for i in tab} >= {
        ("MISSING_COLUMN", "instagram"),
        ("DUPLICATE_COLUMN", "phone"),
        ("CITY_NOT_TAB", "city"),
    }
    assert ("MISSING_COLUMN", "city") not in {(i.code, i.field) for i in tab}


def test_one_click_fixes() -> None:
    values = CLEAN | {"website": "https://facebook.com/mc", "facebook": ""}
    assert fix_for("WEBSITE_IS_FACEBOOK", values) == {
        "label": "Move it to Facebook",
        "changes": {"website": "", "facebook": "https://facebook.com/mc"},
    }
    filled = fix_for("WEBSITE_IS_FACEBOOK", CLEAN | {"website": "https://facebook.com/mc"})
    assert filled == {"label": "Clear Website (Facebook is filled)", "changes": {"website": ""}}
    swapped = fix_for("REVIEWS_RATING_SWAPPED", CLEAN | {"reviews": "4.8", "rating": "120"})
    assert swapped is not None and swapped["changes"] == {"reviews": "120", "rating": "4.8"}
    moved = fix_for(
        "ADDRESS_HAS_PHONE",
        CLEAN | {"address": "5874 S 129th E Ave, (918) 555-0101, Tulsa", "phone": ""},
    )
    assert moved == {
        "label": "Move the phone number to Phone",
        "changes": {"address": "5874 S 129th E Ave, Tulsa", "phone": "(918) 555-0101"},
    }
    assert fix_for("PHONE_WRONG_LENGTH", CLEAN) is None


# --- in the Hub -----------------------------------------------------------------------------------


@pytest.fixture
def owner(client: TestClient, make_user: UserFactory) -> User:
    user = make_user(Role.OWNER, name="Olivia")
    sign_in(client, user.email)
    return user


def bad_row() -> list[Any]:
    """Sparkle's row with a phone number in its address and a Facebook link as website."""
    values = row("Sparkle", "https://share.google/sp", "918 555 0300", website="https://fb.com/sp")
    values[STANDARD.index("Address")] = "9 Elm St (918) 555-0300"
    return values


FILE = {
    "Tulsa": [
        [*STANDARD, "Phone"],  # Phone twice
        [*row("Multi-Clean", "https://share.google/mc", "(918) 523-8300"), ""],
        [*row("Top Shine", "https://share.google/ts", "(918) 555-0000"), ""],
        [*bad_row(), ""],
    ],
    "Dallas": [
        [h for h in STANDARD if h != "Instagram"],  # no Instagram column
        [
            v
            for h, v in zip(
                STANDARD,
                row("Maid Pro", "https://share.google/mp", "214 555 0100", city="Dallas"),
                strict=True,
            )
            if h != "Instagram"
        ],
    ],
}


def save(client: TestClient, content: bytes, name: str = "leads.xlsx") -> dict[str, Any]:
    preview = upload(client, content, name=name)
    saved = client.post(f"/api/imports/{preview['id']}/confirm")
    assert saved.status_code == 200, saved.text
    body: dict[str, Any] = saved.json()
    return body | {"preview_check": preview["preview"]["check"]}


def test_the_preview_shows_lead_check_before_saving(client: TestClient, owner: User) -> None:
    preview = upload(client, workbook(FILE))

    found = preview["preview"]["check"]
    assert found["errors"] == 2  # Sparkle: phone in address, Facebook link as website
    assert found["tabs"]["Tulsa"]["errors"] == 2
    columns = {(c["code"], c["field"]) for c in found["tabs"]["Tulsa"]["columns"]}
    assert ("DUPLICATE_COLUMN", "phone") in columns
    dallas = {(c["code"], c["field"]) for c in found["tabs"]["Dallas"]["columns"]}
    assert ("MISSING_COLUMN", "instagram") in dallas
    rows = client.get(f"/api/imports/{preview['id']}/sheet", params={"tab": "Tulsa"}).json()
    sparkle = next(r for r in rows["rows"] if r["values"]["business_name"] == "Sparkle")
    assert {i["code"] for i in sparkle["issues"]} == {"ADDRESS_HAS_PHONE", "WEBSITE_IS_FACEBOOK"}


def test_a_file_with_an_error_holds_all_its_leads_until_fixed(
    client: TestClient, owner: User, db: Session
) -> None:
    saved = save(client, workbook(FILE), "texas.xlsx")
    batch = batch_id(client, "Batch 1")
    token = bot(client, "gbp-1", "gbp-audit")
    put_bot(client, batch, "gbp-1", "gbp-audit")
    for city, lead in (("Tulsa", "L-000002"), ("Dallas", "")):
        if lead:  # a Top chosen: Tulsa is ready; Dallas has none and stays not ready
            body = {"batch_id": batch, "city": city, "lead_ids": [lead], "top": True}
            assert client.post("/api/leads/top", json=body).status_code == 200

    client.put(f"/api/batches/{batch}/cities", json={"active": True})  # small cities on

    assert saved["check_errors"] == 1
    assert ready(client, token, batch) == []  # Multi-Clean is fine, but its file is held
    alert = db.scalar(select(Alert).where(Alert.kind == "lead_check"))
    assert alert is not None and alert.title == "Lead Check: 2 errors in texas.xlsx"
    assert "Tulsa 2" in alert.body and "agents wait" in alert.body
    cities = {c["city"]: c for c in client.get(f"/api/batches/{batch}/cities").json()}
    assert (cities["Tulsa"]["errors"], cities["Dallas"]["errors"]) == (2, 0)

    sparkle = client.get("/api/leads", params={"batch_id": batch, "problems": "errors"}).json()
    assert [x["business_name"] for x in sparkle["leads"]] == ["Sparkle"]
    lead_id = sparkle["leads"][0]["lead_id"]
    one = client.get("/api/leads", params={"batch_id": batch, "problems": "WEBSITE_IS_FACEBOOK"})
    assert [x["lead_id"] for x in one.json()["leads"]] == [lead_id]
    detail = client.get(f"/api/leads/{lead_id}").json()
    assert detail["held"] is True
    fix = next(i["fix"] for i in detail["check"] if i["code"] == "WEBSITE_IS_FACEBOOK")

    # Fixed in the Hub: the one-click fix, then the address by hand.
    client.patch(f"/api/leads/{lead_id}", json=fix["changes"])
    still = client.get(f"/api/imports/{saved['id']}").json()
    assert still["check_errors"] == 1
    fixed = client.patch(f"/api/leads/{lead_id}", json={"address": "9 Elm St, Tulsa, OK"})
    assert fixed.json()["check"] == [] and fixed.json()["held"] is False
    assert fixed.json()["data"]["facebook"] == "https://fb.com/sp"

    assert client.get(f"/api/imports/{saved['id']}").json()["check_errors"] == 0
    assert ready(client, token, batch) == ["L-000001", lead_id]


def test_the_report_and_its_excel(client: TestClient, owner: User) -> None:
    saved = save(client, workbook(FILE))

    report = client.get(f"/api/imports/{saved['id']}/check").json()

    assert (report["held"], report["errors"]) == (True, 2)
    assert [(c["city"], c["errors"]) for c in report["cities"]] == [("Tulsa", 2)]
    first = report["issues"][0]
    assert (first["business_name"], first["severity"]) == ("Sparkle", "error")
    assert ("DUPLICATE_COLUMN", "phone") in {
        (i["code"], i["field"]) for i in report["tabs"]["Tulsa"]
    }
    excel = client.get(f"/api/imports/{saved['id']}/check.xlsx")
    assert excel.status_code == 200
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert sheet is not None
    lines = [[c.value for c in line] for line in sheet.iter_rows(min_row=2)]
    assert lines[0][:2] == ["Error", "Tulsa"]
    assert any(line[4] == "Instagram" and line[1] == "Dallas" for line in lines)
    again = client.post(f"/api/imports/{saved['id']}/check")
    assert again.json()["errors"] == 2


def test_a_clean_file_isnt_held_by_another(client: TestClient, owner: User, db: Session) -> None:
    held = save(client, workbook(FILE), "bad.xlsx")
    clean = save(
        client,
        workbook(
            {
                "Austin": [
                    STANDARD,
                    row("Shine ATX", "https://share.google/atx", "512 555 0100", city="Austin"),
                ]
            }
        ),
        "good.xlsx",
    )

    assert (held["check_errors"], clean["check_errors"]) == (1, 0)
    titles = [a.title for a in db.scalars(select(Alert).where(Alert.kind == "lead_check"))]
    assert "Lead Check: good.xlsx is clear" in titles


def test_every_field_can_be_edited_but_not_top_and_not_by_viewers(
    client: TestClient, owner: User, make_user: UserFactory
) -> None:
    save(client, workbook(FILE))
    batch = batch_id(client, "Batch 1")
    body = {
        "business_name": "Multi Clean LLC",
        "gbp_link": "https://maps.google.com/?cid=9",
        "address": "1 Main St, Broken Arrow, OK",
        "city": "Broken Arrow",
        "whatsapp": "918 555 0909",
        "facebook": "https://facebook.com/mc",
        "instagram": "https://instagram.com/mc",
        "reviews": "15",
        "rating": "4.9",
    }

    changed = client.patch("/api/leads/L-000001", json=body).json()

    assert (changed["business_name"], changed["city"], changed["gbp_link"]) == (
        "Multi Clean LLC",
        "Broken Arrow",
        "https://maps.google.com/?cid=9",
    )
    assert changed["data"]["rating"] == "4.9"
    cities = [c["city"] for c in client.get(f"/api/batches/{batch}/cities").json()]
    assert "Broken Arrow" in cities
    assert client.patch("/api/leads/L-000001", json={"level": "Top"}).status_code == 422
    assert client.patch("/api/leads/L-000001", json={"city": ""}).status_code == 422
    client.cookies.clear()
    sign_in(client, make_user(Role.VIEWER).email)
    assert client.patch("/api/leads/L-000001", json={"rating": "5"}).status_code == 403
