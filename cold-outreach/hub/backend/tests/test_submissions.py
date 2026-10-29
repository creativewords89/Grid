"""The Files screen (SPEC steps 31 and 39): Collectors hand in lead files and fix them until
Lead Check finds no error; sent on to Top Review, a file is theirs no more. A business
handed in twice is an error, a city handed in twice is muted."""

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.comments import Comment

from app.db.models import Role
from app.imports import submissions
from tests.conftest import Person, run_file_checker
from tests.hub_data import make_benchmark
from tests.test_imports import STANDARD, row, workbook

pytestmark = pytest.mark.usefixtures("files_checked")

GOOD = {
    "Tulsa": [
        STANDARD,
        row("Multi-Clean", "https://share.google/mc", "(918) 523-8300", "kevin@mc.com"),
        row("Top Shine", "https://share.google/ts", "(918) 555-0000"),
    ],
    "Dallas": [STANDARD, row("Sparkle", "https://share.google/sp", "214 555 0100", city="Dallas")],
}
BAD_ROW = row("Bad Co", "https://share.google/bc", "call us", website="https://facebook.com/bc")
BAD = {"Tulsa": [*GOOD["Tulsa"], BAD_ROW], "Dallas": GOOD["Dallas"]}


# --- Lead Check on a whole file -----------------------------------------------------------------


def test_every_problem_has_its_tab_row_and_column() -> None:
    found = submissions.run(workbook(BAD), "Cleaning")
    assert found.businesses == 4
    assert found.check["errors"] == 2
    assert found.check["tabs"]["Tulsa"]["errors"] == 2
    assert found.check["tabs"]["Dallas"]["errors"] == 0
    assert {(p["tab"], p["line"], p["field"], p["code"]) for p in found.problems} == {
        ("Tulsa", 4, "phone", "PHONE_HAS_LETTERS"),
        ("Tulsa", 4, "website", "WEBSITE_IS_FACEBOOK"),
    }
    assert found.problems[0]["business"] == "Bad Co"
    assert submissions.run(workbook(GOOD), "Cleaning").check["errors"] == 0


def test_a_business_on_two_rows_is_checked_once_and_is_an_error() -> None:
    twice = {"Tulsa": [STANDARD, BAD_ROW, BAD_ROW]}
    found = submissions.run(workbook(twice), "Cleaning")
    assert found.check["errors"] == 3
    same = next(p for p in found.problems if p["code"] == "SAME_IN_FILE")
    assert same["message"] == (
        "Same business on another row of this file (Tulsa row 2, Tulsa row 3):"
        " Upload fixes keeps it once"
    )


def test_a_muted_city_is_not_checked() -> None:
    others = submissions.Others(cities={"tulsa": "Tulsa is already in file F1 (a.xlsx)"})
    found = submissions.run(workbook(BAD), "Cleaning", others=others)
    assert found.check["errors"] == 0
    assert found.check["cities"]["tulsa"]["muted"] == "Tulsa is already in file F1 (a.xlsx)"
    assert found.check["cities"]["dallas"]["muted"] == ""
    assert all(lead.get("muted") for lead in found.leads if lead["tab"] == "Tulsa")


def test_a_business_handed_in_before_is_an_error() -> None:
    def same(values: dict[str, str]) -> str | None:
        return "in file F1 (a.xlsx) as F1-0001" if values["business_name"] == "Sparkle" else None

    found = submissions.run(workbook(GOOD), "Cleaning", others=submissions.Others(same=same))
    assert [(p["business"], p["code"], p["message"]) for p in found.problems] == [
        (
            "Sparkle",
            "SAME_ELSEWHERE",
            "Same business already handed in: in file F1 (a.xlsx) as F1-0001",
        )
    ]


def test_the_same_leads_have_the_same_fingerprint() -> None:
    first = submissions.run(workbook(GOOD), "Cleaning").fingerprint
    resaved = load_workbook(io.BytesIO(workbook(GOOD)))
    resaved["Tulsa"]["Z1"] = None
    out = io.BytesIO()
    resaved.save(out)
    assert submissions.run(out.getvalue(), "Cleaning").fingerprint == first
    assert submissions.run(workbook(BAD), "Cleaning").fingerprint != first


def test_marks_of_an_old_marked_copy_are_taken_out() -> None:
    """Files handed in before step 34 were downloaded with a Lead Check tab and notes."""
    book = load_workbook(io.BytesIO(workbook(BAD)))
    report = book.create_sheet(submissions.REPORT_TAB, 0)
    report.append(["Severity", "Tab", "Row"])
    phone = book["Tulsa"].cell(4, STANDARD.index("Phone") + 1)
    phone.value = "(918) 555-2222"
    phone.fill = submissions.RED
    phone.comment = Comment("Error: Phone holds letters", submissions.AUTHOR)
    book["Tulsa"].cell(4, STANDARD.index("Website") + 1).value = "https://badco.com"
    out = io.BytesIO()
    book.save(out)
    cleaned = load_workbook(io.BytesIO(submissions.clean(out.getvalue())))
    assert cleaned.sheetnames == ["Tulsa", "Dallas"]
    assert cleaned["Tulsa"].cell(4, STANDARD.index("Phone") + 1).comment is None
    assert submissions.run(submissions.clean(out.getvalue()), "Cleaning").check["errors"] == 0


# --- the API ------------------------------------------------------------------------------------


@pytest.fixture
def owner(person: Person) -> TestClient:
    client, _ = person(Role.OWNER, "Olivia")
    make_benchmark(client, "Cleaning Benchmark 1", "Cleaning")
    make_benchmark(client, "Roofing Benchmark 1", "Roofing")
    return client


def hand_in(
    client: TestClient,
    tabs: dict[str, Any],
    niche: str = "Cleaning",
    name: str = "tulsa.xlsx",
    send: bool = True,
) -> dict[str, Any]:
    response = client.post(
        "/api/submissions", files={"file": (name, workbook(tabs))}, data={"niche": niche}
    )
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "checking"
    run_file_checker()
    file_id = response.json()["id"]
    body: dict[str, Any] = client.get(f"/api/submissions/{file_id}").json()
    if send and body["status"] == "final_check":
        body = sent_on(client, file_id)
    return body


def sent_on(client: TestClient, file_id: str) -> dict[str, Any]:
    """Send a file without errors on to Top Review, as its sender does (step 36)."""
    response = client.post(f"/api/submissions/{file_id}/send")
    assert response.status_code == 200, response.text
    return client.get(f"/api/submissions/{file_id}").json()  # type: ignore[no-any-return]


def revise(client: TestClient, file_id: str, tabs: dict[str, Any], send: bool = True) -> Any:
    response = client.post(
        f"/api/submissions/{file_id}/version", files={"file": ("tulsa-v2.xlsx", workbook(tabs))}
    )
    if response.status_code != 200:
        return response
    run_file_checker()
    if send and client.get(f"/api/submissions/{file_id}").json()["status"] == "final_check":
        client.post(f"/api/submissions/{file_id}/send")
    return client.get(f"/api/submissions/{file_id}")


def test_a_collector_sees_only_the_files_screen(owner: TestClient, person: Person) -> None:
    collector, _ = person(Role.COLLECTOR, "Cal")
    for path in (
        "/api/batches",
        "/api/leads",
        "/api/imports",
        "/api/niches",
        "/api/agents",
        "/api/runs",
        "/api/benchmarks",
        "/api/alerts",
        "/api/users",
        "/api/review/cities",
    ):
        assert collector.get(path).status_code == 403, path
    assert collector.get("/api/submissions").status_code == 200
    assert collector.get("/api/submissions/niches").json() == ["Cleaning", "Roofing"]
    assert collector.get("/api/auth/me").json()["user"]["role"] == "collector"


def test_a_collector_hands_in_fixes_and_sees_only_their_own_files(
    owner: TestClient, person: Person
) -> None:
    cal, _ = person(Role.COLLECTOR, "Cal")
    dana, _ = person(Role.COLLECTOR, "Dana")
    bad = hand_in(cal, BAD)
    assert bad["status"] == "errors" and bad["errors"] == 2 and bad["version"] == 1
    assert bad["uploaded_by"] == "Cal" and bad["mine"]
    assert bad["batch_name"] == "Cleaning"
    assert len(bad["problems"]) == 2 and bad["tabs"]["Tulsa"]["errors"] == 2
    assert [(c["name"], c["businesses"], c["state"]) for c in bad["cities"]] == [
        ("Dallas", 1, "small"),
        ("Tulsa", 3, "small"),
    ]
    # Others' files are not there for a Collector; Owners see every file.
    assert dana.get("/api/submissions").json() == []
    assert dana.get(f"/api/submissions/{bad['id']}").status_code == 404
    assert dana.get(f"/api/submissions/{bad['id']}/marked.xlsx").status_code == 404
    assert revise(dana, bad["id"], GOOD).status_code == 404
    assert dana.delete(f"/api/submissions/{bad['id']}").status_code == 404
    assert [f["file_name"] for f in owner.get("/api/submissions").json()] == ["tulsa.xlsx"]
    # The fix file, then a revised version: checked again.
    marked = cal.get(f"/api/submissions/{bad['id']}/marked.xlsx")
    assert marked.status_code == 200
    assert load_workbook(io.BytesIO(marked.content)).sheetnames == ["Problems"]
    fixed = revise(cal, bad["id"], GOOD).json()
    assert fixed["errors"] == 0 and fixed["version"] == 2 and fixed["problems"] == []
    # Sent on, its cities are all under 10 businesses: none goes on unless given a Top.
    assert fixed["status"] == "in_batch" and fixed["cities_going"] == 0


def test_handing_in_needs_a_known_niche_and_an_excel_file(
    owner: TestClient, person: Person
) -> None:
    cal, _ = person(Role.COLLECTOR)
    wrong = cal.post(
        "/api/submissions", files={"file": ("a.xlsx", workbook(GOOD))}, data={"niche": "Bakery"}
    )
    assert wrong.status_code == 422 and wrong.json()["error"]["code"] == "unknown_niche"
    junk = cal.post(
        "/api/submissions", files={"file": ("a.xlsx", b"x")}, data={"niche": "cleaning"}
    )
    # Taken at once and checked in the background: it fails there, and says why.
    assert junk.status_code == 201 and junk.json()["status"] == "checking"
    assert run_file_checker() == 1
    failed = cal.get(f"/api/submissions/{junk.json()['id']}").json()
    assert failed["status"] == "failed" and "isn't an Excel file" in failed["error"]
    assert cal.get(f"/api/submissions/{junk.json()['id']}/marked.xlsx").status_code == 409
    assert revise(cal, junk.json()["id"], GOOD, send=False).json()["status"] == "final_check"
    assert hand_in(cal, BIG_DALLAS, niche="roofing")["niche"] == "Roofing"


def test_viewers_see_every_file_but_cannot_hand_in_or_send(
    owner: TestClient, person: Person
) -> None:
    cal, _ = person(Role.COLLECTOR)
    viewer, _ = person(Role.VIEWER)
    sent = hand_in(cal, GOOD)
    assert [f["id"] for f in viewer.get("/api/submissions").json()] == [sent["id"]]
    refused = viewer.post(
        "/api/submissions", files={"file": ("a.xlsx", workbook(GOOD))}, data={"niche": "Cleaning"}
    )
    assert refused.status_code == 403
    assert viewer.post(f"/api/submissions/{sent['id']}/to-batch", json={}).status_code == 403
    assert viewer.delete(f"/api/submissions/{sent['id']}").status_code == 403
    tulsa = viewer.post(f"/api/submissions/{sent['id']}/cities/tulsa/unmute")
    assert tulsa.status_code == 403


def test_a_file_sent_on_is_locked_for_its_sender(owner: TestClient, person: Person) -> None:
    cal, _ = person(Role.COLLECTOR)
    first = hand_in(cal, BAD)
    assert cal.delete(f"/api/submissions/{first['id']}").status_code == 204
    assert cal.get("/api/submissions").json() == []
    waiting = hand_in(cal, GOOD, send=False)
    assert waiting["status"] == "final_check"
    assert revise(cal, waiting["id"], GOOD, send=False).status_code == 200
    sent = sent_on(cal, waiting["id"])
    # Sent on to Top Review: only seen. An Owner or Admin can still delete it.
    assert revise(cal, sent["id"], GOOD).json()["error"]["code"] == "sent_on"
    assert cal.delete(f"/api/submissions/{sent['id']}").json()["error"]["code"] == "sent_on"
    upload = cal.post(
        f"/api/submissions/{sent['id']}/fixes", files={"file": ("f.xlsx", workbook(GOOD))}
    )
    assert upload.json()["error"]["code"] == "sent_on"
    assert owner.delete(f"/api/submissions/{sent['id']}").status_code == 204


# --- the same business or city twice (step 39) -------------------------------------------------

TWINS = {
    "Austin": [STANDARD, row("Sparkle", "https://share.google/sp", "512 555 0100", city="Austin")]
}
BIG_DALLAS = {
    "Dallas": [
        STANDARD,
        *(
            row(f"Gleam {n}", f"https://share.google/g{n}", f"(214) 555-{n:04d}", city="Dallas")
            for n in range(10)
        ),
    ]
}


def test_a_business_already_handed_in_is_an_error(owner: TestClient, person: Person) -> None:
    cal, _ = person(Role.COLLECTOR)
    first = hand_in(cal, GOOD)
    # Sparkle (by its GBP link) is in the earlier file, in another city and another niche too.
    second = hand_in(cal, TWINS, niche="Roofing", name="austin.xlsx")
    assert second["status"] == "errors" and second["errors"] == 1
    assert second["problems"][0]["message"] == (
        f"Same business already handed in: in file F{first['number']} (tulsa.xlsx) as "
        f"F{first['number']}-0003"
    )
    # The earlier file isn't told about the later one.
    assert owner.get(f"/api/submissions/{first['id']}").json()["errors"] == 0
    # Taken out of the second file, it passes.
    fixed = revise(
        cal,
        second["id"],
        {
            "Austin": [
                STANDARD,
                row("New Co", "https://share.google/n", "512 555 0199", city="Austin"),
            ]
        },
    )
    assert fixed.json()["errors"] == 0


def test_a_city_handed_in_again_is_muted_until_unmuted(owner: TestClient, person: Person) -> None:
    cal, _ = person(Role.COLLECTOR)
    first = hand_in(cal, BIG_DALLAS, name="dallas.xlsx")
    other_dallas = {
        "Dallas": [
            STANDARD,
            *(
                row(f"Shine {n}", f"https://share.google/s{n}", f"(214) 556-{n:04d}", city="Dallas")
                for n in range(10)
            ),
            row(
                "Bad Co",
                "https://share.google/bc",
                "call us",
                city="Dallas",
                website="https://facebook.com/bc",
            ),
        ],
        "Austin": [
            STANDARD,
            row("Bloom", "https://share.google/bl", "512 555 0111", city="Austin"),
        ],
    }
    second = hand_in(cal, other_dallas, name="dallas-2.xlsx")
    # Dallas is muted (its bad row isn't counted); Austin is checked as usual.
    assert second["errors"] == 0 and second["cities_muted"] == 1
    dallas = next(c for c in second["cities"] if c["name"] == "Dallas")
    assert dallas["state"] == "muted"
    assert dallas["muted"] == (
        f"Dallas is already in file F{first['number']} (dallas.xlsx), handed in earlier"
    )
    # The same city of another niche isn't muted.
    roofing = hand_in(cal, other_dallas, niche="Roofing", name="roof.xlsx")
    assert roofing["cities_muted"] == 0
    # Only Owners and Admins unmute, and only once no other file has it.
    url = f"/api/submissions/{second['id']}/cities/dallas/unmute"
    assert cal.post(url).status_code == 403
    still = owner.post(url)
    assert still.json()["error"]["code"] == "still_taken"
    assert owner.delete(f"/api/submissions/{first['id']}").status_code == 204
    unmuted = owner.post(url)
    assert unmuted.status_code == 200 and unmuted.json()["status"] == "checking"
    run_file_checker()
    again = owner.get(f"/api/submissions/{second['id']}").json()
    # Checked again: Dallas is on, and its bad row (and Roofing's twins) count now.
    assert again["cities_muted"] == 0
    assert again["status"] == "errors"


def test_owners_invite_collectors(owner: TestClient) -> None:
    response = owner.post(
        "/api/users/invite", json={"name": "Cal", "email": "cal@example.com", "role": "collector"}
    )
    assert response.status_code == 201, response.text
    assert response.json()["user"]["role"] == "collector"
