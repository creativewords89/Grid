"""Importing lead files (SPEC section 6.4), from fixtures built like the team's real files."""

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import AuditEntry, Batch, DoNotContact, Import, Lead, LeadList, Role, User
from app.imports import reader
from tests.conftest import UserFactory, sign_in
from tests.hub_data import batch_id

STANDARD = [
    "Business name",
    "MAP",
    "Address",
    "City",
    "Website",
    "Phone",
    "WA",
    "Facebook",
    "Instagram",
    "Email",
    "Reviews",
    "Average",
]


def workbook(tabs: dict[str, list[list[Any]]]) -> bytes:
    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    for title, rows in tabs.items():
        sheet = book.create_sheet(title)
        for line in rows:
            sheet.append(line)
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # text that looks like a formula, as people type it
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def row(
    name: str, link: str, phone: str, email: str = "", city: str = "Tulsa", **extra: Any
) -> list[Any]:
    values = {
        "Business name": name,
        "MAP": link,
        "Address": f"1 {name} St, {city}",
        "City": city,
        "Website": extra.get("website", f"https://{name.lower().replace(' ', '')}.com/"),
        "Phone": phone,
        "WA": "",
        "Facebook": "",
        "Instagram": "",
        "Email": email,
        "Reviews": 12,
        "Average": 4.5,
    }
    return [values[h] for h in STANDARD]


# --- reading -----------------------------------------------------------------------------------


def test_variant_headers_and_a_missing_city_column() -> None:
    content = workbook(
        {
            "Dallas": [
                [
                    "Business Name ",
                    "Map URl",
                    "Address ",
                    "Website ",
                    "Phone ",
                    "Whatsapp",
                    "Facebook",
                    "instagram",
                    "Gmail",
                    "Number of Reviews",
                    "Average Ratings",
                    "Comment",
                    "Operation Hour",
                ],
                [
                    "Sparkle Co\n",
                    "https://share.google/abc",
                    " 1 Main St ",
                    "https://sparkle.com/?utm_source=gmb&utm_medium=x&page=2",
                    "(214) 555-0100",
                    "Yes",
                    "",
                    "",
                    "INFO@Sparkle.com\n",
                    33,
                    4.9,
                    "call after 5",
                    "9-5",
                ],
            ]
        }
    )

    report = reader.read(content, "Cleaning")

    [lead] = report.leads
    assert lead.values["city"] == "Dallas"  # from the tab name
    assert lead.values["business_name"] == "Sparkle Co"
    assert lead.values["website"] == "https://sparkle.com/?page=2"
    assert lead.values["whatsapp"] == "Yes"
    assert (lead.values["reviews"], lead.values["rating"]) == ("33", "4.9")
    assert lead.values["notes"] == "call after 5"
    assert lead.emails == ["info@sparkle.com"]
    assert report.unknown_columns == {"Operation Hour": ["Dallas"]}


def test_same_name_and_address_and_two_emails_become_one_lead() -> None:
    content = workbook(
        {
            "Aurora": [
                STANDARD,
                row("Four Star Drywall", "", "(303) 238-5594", "contactus@fourstar.com"),
                row("Four Star Drywall", "", "(303) 238-5594", "Tony.jr@fourstar.com"),
            ]
        }
    )

    report = reader.read(content, "Construction")

    [lead] = report.leads
    assert lead.emails == ["contactus@fourstar.com", "tony.jr@fourstar.com"]
    assert report.merged_rows == [["Aurora row 2", "Aurora row 3"]]


def test_one_phone_but_a_profile_per_city_stays_separate() -> None:
    content = workbook(
        {
            "Knoxville": [STANDARD, row("MaidPro", "https://share.google/knox", "(865) 555-0100")],
            "Chattanooga": [
                STANDARD,
                row("MaidPro", "https://share.google/chatt", "(865) 555-0100", city="Chattanooga"),
            ],
        }
    )

    assert len(reader.read(content, "Cleaning").leads) == 2


def test_city_filled_down_an_empty_tab_is_not_a_business() -> None:
    content = workbook(
        {
            "Jackson": [STANDARD, *([["", "", "", "Jackson"]] * 50)],
            "Tulsa": [STANDARD, row("Level Up", "https://share.google/lvl", "918 555 0101")],
        }
    )

    report = reader.read(content, "Cleaning")

    assert len(report.leads) == 1
    assert report.empty_tabs == ["Jackson"]
    assert report.skipped_rows == []


def test_a_name_without_an_address_is_never_matched() -> None:
    """Step 39: only the GBP link, or the name and the address, make two rows one."""
    content = workbook(
        {
            "San Jose": [
                ["Business name", "City", "Email"],
                ["MOLLY MAID of S. Silicon Valley", "San Jose", "a@mm.com"],
                ["Molly Maid of S Silicon Valley", "San Jose", "b@mm.com"],
            ]
        }
    )

    assert len(reader.read(content, "Cleaning").leads) == 2


def test_one_brand_in_many_cities_or_one_phone_stays_separate() -> None:
    """Step 39: the same name at another address, or the same phone, is another business."""
    content = workbook(
        {
            "Austin": [STANDARD, row("Molly Maid", "", "(512) 555-0100", city="Austin")],
            "Dallas": [STANDARD, row("Molly Maid", "", "(512) 555-0100", city="Dallas")],
        }
    )

    assert len(reader.read(content, "Cleaning").leads) == 2


def test_bad_emails_are_listed_and_encoded_spaces_cleaned() -> None:
    content = workbook(
        {"Tulsa": [STANDARD, row("A", "https://share.google/a", "1", "support@a.com%20, http:")]}
    )

    report = reader.read(content, "Cleaning")

    assert report.leads[0].emails == ["support@a.com"]
    assert report.invalid_emails == ["Tulsa row 2: http:"]


def test_not_an_excel_file() -> None:
    with pytest.raises(reader.ReadError, match="isn't an Excel file"):
        reader.read(b"name,phone\nA,1\n", "Cleaning")


# --- the API -----------------------------------------------------------------------------------


@pytest.fixture
def owner(client: TestClient, make_user: UserFactory) -> User:
    user = make_user(Role.OWNER, name="Olivia")
    sign_in(client, user.email)
    return user


def leads(client: TestClient, **params: Any) -> dict[str, dict[str, Any]]:
    """The leads the Leads screen shows, by Lead ID."""
    page = client.get("/api/leads", params=params).json()
    return {lead["lead_id"]: lead for lead in page["leads"]}


def mark_top(client: TestClient, *lead_ids: str, top: bool = True) -> None:
    """Marks leads Top in their city, as the Cities tab does (all in one city)."""
    first = client.get(f"/api/leads/{lead_ids[0]}").json()
    body = {"batch_id": first["batch_id"], "city": first["city"], "lead_ids": list(lead_ids)}
    response = client.post("/api/leads/top", json={**body, "top": top})
    assert response.status_code == 200, response.text


def upload(
    client: TestClient, content: bytes, batch: str = "Batch 1", name: str = "leads.xlsx"
) -> dict[str, Any]:
    response = client.post(
        "/api/imports",
        files={"file": (name, content)},
        data={"batch_id": batch_id(client, batch)},
    )
    assert response.status_code == 201, response.text
    body: dict[str, Any] = response.json()
    return body


FILE = {
    "Tulsa": [
        [*STANDARD, "Level"],
        [*row("Multi-Clean", "https://share.google/mc", "(918) 523-8300", "kevin@mc.com"), ""],
        [*row("Top Shine", "https://share.google/ts", "(918) 555-0000", "boss@ts.com"), "Top"],
        [*row("No Map Co", "", "(918) 555-1111", '=HYPERLINK("x")'), ""],
    ],
    "Dallas": [STANDARD, row("Sparkle", "https://share.google/sp", "214 555 0100", city="Dallas")],
}


def test_preview_saves_nothing_but_the_import(client: TestClient, owner: User, db: Session) -> None:
    body = upload(client, workbook(FILE))

    assert body["status"] == "preview"
    counts = body["counts"]
    assert (counts["leads"], counts["new"], counts["top"], counts["no_gbp_link"]) == (4, 4, 0, 1)
    assert counts["invalid_emails"] == 1
    assert sorted(body["preview"]["new_lists"]) == ["Dallas", "Tulsa"]
    # Top is chosen in the Hub: the file's Level column is only another column.
    assert body["preview"]["has_level_column"] is True
    assert body["preview"]["top"] == []
    assert db.scalar(select(Lead.lead_id)) is None


def test_confirm_gives_ids_keeps_the_rows_and_protects_top_leads(
    client: TestClient, owner: User, db: Session
) -> None:
    preview = upload(client, workbook(FILE))

    done = client.post(f"/api/imports/{preview['id']}/confirm")

    assert done.status_code == 200, done.text
    assert (done.json()["counts"]["saved_new"], done.json()["counts"]["saved_updated"]) == (4, 0)
    saved = leads(client)
    assert list(saved) == ["L-000001", "L-000002", "L-000003", "L-000004"]
    first = saved["L-000001"]
    assert (first["city"], first["niche"], first["batch_name"], first["business_name"]) == (
        "Tulsa",
        "Cleaning",
        "Batch 1",
        "Multi-Clean",
    )
    assert first["emails"] == ["kevin@mc.com"]
    assert first["data"]["phone"] == "(918) 523-8300"
    assert first["data"]["website"] == "https://multi-clean.com/"
    assert saved["L-000002"]["level"] == ""  # the file's "Top" changes nothing
    assert saved["L-000002"]["extra"] == {"Level": "Top"}
    assert db.scalar(select(DoNotContact)) is None
    mark_top(client, "L-000002")
    blocked = {(d.kind, d.value) for d in db.scalars(select(DoNotContact))}
    assert blocked == {
        ("gbp", "https://share.google/ts"),
        ("phone", "9185550000"),
        ("email", "boss@ts.com"),
    }
    lists = {(x.city, x.niche, str(x.batch_id)) for x in db.scalars(select(LeadList))}
    assert lists == {
        ("Tulsa", "Cleaning", preview["batch_id"]),
        ("Dallas", "Cleaning", preview["batch_id"]),
    }
    assert db.get(Import, preview["id"]).rows == []  # type: ignore[union-attr]


def test_a_second_import_updates_instead_of_adding(client: TestClient, owner: User) -> None:
    first = upload(client, workbook(FILE))
    client.post(f"/api/imports/{first['id']}/confirm")
    mark_top(client, "L-000002")
    again = workbook(
        {
            "Tulsa": [
                STANDARD,
                row(
                    "Multi-Clean",
                    "https://share.google/mc",
                    "(918) 523-8300",
                    "new@mc.com",
                    website="",
                ),
                row("Brand New", "https://share.google/bn", "(918) 555-2222"),
            ]
        }
    )

    preview = upload(client, again)
    assert (preview["counts"]["new"], preview["counts"]["updated"]) == (1, 1)
    client.post(f"/api/imports/{preview['id']}/confirm")

    saved = leads(client)
    assert len(saved) == 5
    multi = saved["L-000001"]
    assert multi["emails"] == ["kevin@mc.com", "new@mc.com"]
    assert multi["data"]["website"] == "https://multi-clean.com/"  # an empty cell never blanks it
    assert saved["L-000002"]["level"] == "Top"  # chosen in the Hub: files don't change it
    assert saved["L-000005"]["business_name"] == "Brand New"


def test_discard_and_confirm_only_once(client: TestClient, owner: User) -> None:
    preview = upload(client, workbook(FILE))

    assert client.post(f"/api/imports/{preview['id']}/discard").json()["status"] == "discarded"
    assert client.post(f"/api/imports/{preview['id']}/confirm").status_code == 409


def test_bad_uploads(client: TestClient, owner: User) -> None:
    batch = batch_id(client)
    not_excel = client.post(
        "/api/imports", files={"file": ("a.csv", b"a,b\n1,2")}, data={"batch_id": batch}
    )
    no_batch = client.post("/api/imports", files={"file": ("a.xlsx", workbook(FILE))})
    unknown = client.post(
        "/api/imports",
        files={"file": ("a.xlsx", workbook(FILE))},
        data={"batch_id": "00000000-0000-0000-0000-000000000000"},
    )

    assert (not_excel.status_code, not_excel.json()["error"]["code"]) == (422, "unreadable")
    assert no_batch.status_code == 422
    assert unknown.status_code == 404


def test_a_business_stays_in_the_batch_it_came_in(client: TestClient, owner: User) -> None:
    first = upload(client, workbook(FILE))
    client.post(f"/api/imports/{first['id']}/confirm")
    again = workbook(
        {
            "Tulsa": [
                STANDARD,
                row("Multi-Clean", "https://share.google/mc", "(918) 523-8300", "new@mc.com"),
                row("Brand New", "https://share.google/bn", "(918) 555-2222"),
            ]
        }
    )

    preview = upload(client, again, batch="Batch 2")
    counts = preview["counts"]
    assert (counts["new"], counts["updated"], counts["other_batch"]) == (1, 0, 1)
    assert preview["preview"]["other_batch"] == ["Multi-Clean (L-000001, in Batch 1)"]
    client.post(f"/api/imports/{preview['id']}/confirm")

    saved = leads(client)
    assert saved["L-000001"]["batch_name"] == "Batch 1"
    assert saved["L-000001"]["emails"] == ["kevin@mc.com", "new@mc.com"]  # still updated
    assert saved["L-000005"]["batch_name"] == "Batch 2"
    batches = {b["name"]: b for b in client.get("/api/batches").json()}
    assert (batches["Batch 1"]["leads"], batches["Batch 2"]["leads"]) == (4, 1)
    assert batches["Batch 2"]["files"] == 1
    mine = client.get("/api/imports", params={"batch_id": batches["Batch 2"]["id"]}).json()
    assert [(i["file_name"], i["batch_name"]) for i in mine] == [("leads.xlsx", "Batch 2")]


@pytest.mark.parametrize("role", [Role.VIEWER])
def test_viewers_see_but_cannot_import(
    client: TestClient, owner: User, make_user: UserFactory, role: Role
) -> None:
    preview = upload(client, workbook(FILE))
    client.cookies.clear()
    sign_in(client, make_user(role).email)

    assert client.get("/api/imports").status_code == 200
    assert client.get(f"/api/imports/{preview['id']}").status_code == 200
    assert client.post(f"/api/imports/{preview['id']}/confirm").status_code == 403
    denied = client.post(
        "/api/imports",
        files={"file": ("a.xlsx", workbook(FILE))},
        data={"batch_id": preview["batch_id"]},
    )
    assert denied.status_code == 403


def test_a_batchs_cities_count_leads_and_can_be_switched_off(
    client: TestClient, owner: User
) -> None:
    preview = upload(client, workbook(FILE))
    client.post(f"/api/imports/{preview['id']}/confirm")
    mark_top(client, "L-000002")
    url = f"/api/batches/{preview['batch_id']}/cities"

    cities = {x["city"]: x for x in client.get(url).json()}
    tulsa = cities["Tulsa"]
    changed = client.patch(f"/api/lists/{tulsa['id']}", json={"active": False, "daily_limit": 20})

    assert (tulsa["leads"], tulsa["top"]) == (3, 1)
    assert changed.json() | {"active": False, "daily_limit": 20} == changed.json()
    off = client.put(url, json={"active": False}).json()
    assert [x["active"] for x in off] == [False, False]
    assert all(x["active"] for x in client.put(url, json={"active": True}).json())


def test_each_city_says_the_file_it_first_came_with(client: TestClient, owner: User) -> None:
    """Step 22: a city shows the file it first came with; the latest file's cities are new;
    deleting that file moves the city to the next file that has its businesses."""
    first = upload(client, workbook(FILE), name="october.xlsx")
    client.post(f"/api/imports/{first['id']}/confirm")
    second = upload(
        client,
        workbook(
            {
                "More": [
                    STANDARD,
                    row("Fresh Tulsa", "https://share.google/ft", "918 555 0300"),
                    row("Shine ATX", "https://share.google/atx", "512 555 0100", city="Austin"),
                ]
            }
        ),
        name="november.xlsx",
    )
    client.post(f"/api/imports/{second['id']}/confirm")
    url = f"/api/batches/{first['batch_id']}/cities"

    def origins() -> dict[str, tuple[str | None, bool]]:
        return {
            x["city"]: ((x["added_from"] or {}).get("file_name"), x["new"])
            for x in client.get(url).json()
        }

    assert origins() == {
        "Austin": ("november.xlsx", True),
        "Dallas": ("october.xlsx", False),
        "Tulsa": ("october.xlsx", False),
    }

    assert client.delete(f"/api/imports/{first['id']}").status_code == 204

    assert origins() == {"Austin": ("november.xlsx", True), "Tulsa": ("november.xlsx", True)}


def test_the_same_city_in_two_batches_is_switched_separately(
    client: TestClient, owner: User
) -> None:
    for batch, content in (
        ("Batch 1", {"Tulsa": [STANDARD, row("A", "https://share.google/a", "1")]}),
        ("Batch 2", {"Tulsa": [STANDARD, row("B", "https://share.google/b", "2")]}),
    ):
        client.post(f"/api/imports/{upload(client, workbook(content), batch)['id']}/confirm")
    first, second = batch_id(client, "Batch 1"), batch_id(client, "Batch 2")
    tulsa = client.get(f"/api/batches/{first}/cities").json()[0]

    assert tulsa["active"] is False  # one business: a small city, muted
    client.patch(f"/api/lists/{tulsa['id']}", json={"active": True})

    assert client.get(f"/api/batches/{first}/cities").json()[0]["active"] is True
    assert client.get(f"/api/batches/{second}/cities").json()[0]["active"] is False


# --- the file tab by tab (step 14) ------------------------------------------------------------

SHEET_FILE = {
    "Tulsa": [
        [*STANDARD, "Level", "Operation Hour"],
        [
            *row("Multi-Clean", "https://share.google/mc", "(918) 523-8300", "Kevin@MC.com"),
            "",
            "9-5",
        ],
        [*row("Top Shine", "https://share.google/ts", "(918) 555-0000", "boss@ts.com"), "Top", ""],
        [*row("", "", "(918) 555-2222", "x@y.com"), "", ""],  # no name, no map: skipped
        [*row("Multi Clean 2", "https://share.google/mc", "", "info@mc.com"), "", "8-6"],
    ],
    "Empty": [["Notes"]],
    "Dallas": [STANDARD, row("Sparkle", "https://share.google/sp", "214 555 0100", city="Dallas")],
}


def sheet(client: TestClient, import_id: str, **params: Any) -> dict[str, Any]:
    response = client.get(f"/api/imports/{import_id}/sheet", params=params)
    assert response.status_code == 200, response.text
    body: dict[str, Any] = response.json()
    return body


def test_the_preview_shows_the_file_tab_by_tab(client: TestClient, owner: User) -> None:
    preview = upload(client, workbook(SHEET_FILE))

    shown = sheet(client, preview["id"])

    assert shown["tabs"] == [
        {"name": "Tulsa", "rows": 4},
        {"name": "Empty", "rows": 0},
        {"name": "Dallas", "rows": 1},
    ]
    labels = [c["label"] for c in shown["columns"]]
    assert labels[:3] == ["Business name", "MAP", "Address"]
    assert labels[-1] == "Operation Hour"
    first = shown["rows"][0]
    assert (first["tab"], first["line"], first["status"], first["merged"]) == (
        "Tulsa",
        2,
        "new",
        True,
    )
    assert first["values"]["email"] == "kevin@mc.com"
    assert first["raw"]["email"] == "Kevin@MC.com"  # the original text of a changed cell
    assert first["extra"] == {"Operation Hour": "9-5"}
    assert [r["status"] for r in shown["rows"]] == ["new", "new", "skipped", "new", "new"]
    assert shown["statuses"]["merged"] == 2 and shown["statuses"]["top"] == 0  # not from files
    assert shown["rows"][1]["extra"] == {"Level": "Top"}
    assert sheet(client, preview["id"], status="skipped")["total"] == 1
    assert sheet(client, preview["id"], tab="Dallas")["total"] == 1
    # A business chosen as Top in the Hub shows as Top when a file has it again.
    client.post(f"/api/imports/{preview['id']}/confirm")
    mark_top(client, "L-000002")
    again = upload(client, workbook(SHEET_FILE))
    tulsa_top = sheet(client, again["id"], tab="Tulsa", status="top")
    assert [r["line"] for r in tulsa_top["rows"]] == [3]


def test_other_columns_are_saved_and_shown_and_never_blanked(
    client: TestClient, owner: User
) -> None:
    first = upload(client, workbook(SHEET_FILE))
    client.post(f"/api/imports/{first['id']}/confirm")
    again = {
        "Tulsa": [
            [*STANDARD, "Operation Hour"],
            [*row("Multi-Clean", "https://share.google/mc", "", ""), ""],
        ]
    }
    second = upload(client, workbook(again))
    assert sheet(client, second["id"])["rows"][0]["status"] == "updated"
    client.post(f"/api/imports/{second['id']}/confirm")

    page = client.get("/api/leads").json()
    multi = next(x for x in page["leads"] if x["business_name"] == "Multi-Clean")
    assert multi["extra"] == {"Operation Hour": "9-5"}  # the empty cell didn't blank it
    assert page["extra_columns"] == ["Level", "Operation Hour"]
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(client.get("/api/leads/export").content)).active
    assert book is not None
    assert "Operation Hour" in next(book.iter_rows(values_only=True))


def test_a_saved_file_opens_tab_by_tab_from_its_leads(
    client: TestClient, owner: User, db: Session
) -> None:
    preview = upload(client, workbook(SHEET_FILE))
    client.post(f"/api/imports/{preview['id']}/confirm")
    stored = db.get(Import, preview["id"])
    assert stored is not None and "values" not in stored.sheet["rows"][0]  # only where it went

    shown = sheet(client, preview["id"], tab="Tulsa")
    first = shown["rows"][0]
    assert first["lead_id"] == "L-000001" and first["values"]["business_name"] == "Multi-Clean"
    assert first["extra"] == {"Operation Hour": "9-5"}
    assert shown["rows"][3]["lead_id"] == "L-000001"  # the merged row went to the same lead
    client.delete("/api/leads/L-000002")
    assert sheet(client, preview["id"], tab="Tulsa")["rows"][1]["deleted"] is True


def test_discarding_drops_the_rows_and_old_files_show_counts_only(
    client: TestClient, owner: User, db: Session
) -> None:
    preview = upload(client, workbook(SHEET_FILE))
    client.post(f"/api/imports/{preview['id']}/discard")

    assert sheet(client, preview["id"])["kept"] is False
    assert sheet(client, preview["id"])["rows"] == []


def test_viewers_see_the_file_but_signed_out_people_do_not(
    client: TestClient, owner: User, make_user: UserFactory
) -> None:
    preview = upload(client, workbook(SHEET_FILE))
    viewer = make_user(Role.VIEWER, name="Vic")
    sign_in(client, viewer.email)

    assert client.get(f"/api/imports/{preview['id']}/sheet").status_code == 200
    client.post("/api/auth/logout")
    assert client.get(f"/api/imports/{preview['id']}/sheet").status_code == 401


# --- deleting a file (step 15) ----------------------------------------------------------------


def test_deleting_a_file_takes_the_leads_it_brought_in_first(
    client: TestClient, owner: User, db: Session
) -> None:
    first = upload(client, workbook(FILE))
    client.post(f"/api/imports/{first['id']}/confirm")
    mark_top(client, "L-000002")
    later = {
        "Tulsa": [
            STANDARD,
            row("Multi-Clean", "https://share.google/mc", "", "new@mc.com"),  # known: stays
            row("Fresh Co", "https://share.google/fc", "(918) 555-7777"),  # new: goes
        ],
        "Austin": [STANDARD, row("Austin Co", "https://share.google/ac", "1", city="Austin")],
    }
    second = upload(client, workbook(later))
    client.post(f"/api/imports/{second['id']}/confirm")

    what = client.get(f"/api/imports/{second['id']}/deletion").json()
    assert what == {"leads_going": 2, "leads_staying": 1, "batch_running": False}
    assert client.delete(f"/api/imports/{second['id']}").status_code == 204

    assert sorted(leads(client)) == ["L-000001", "L-000002", "L-000003", "L-000004"]
    cities = [x["city"] for x in client.get(f"/api/batches/{first['batch_id']}/cities").json()]
    assert cities == ["Dallas", "Tulsa"]  # Austin had no leads left
    assert [i["id"] for i in client.get("/api/imports").json()] == [first["id"]]
    entry = db.scalar(
        select(AuditEntry).where(AuditEntry.entity == "import", AuditEntry.action == "delete")
    )
    assert entry is not None and entry.title == "leads.xlsx"
    # The Top lead's do-not-contact entries stay when its file goes.
    assert client.delete(f"/api/imports/{first['id']}").status_code == 204
    assert leads(client) == {}
    assert db.scalar(select(DoNotContact).where(DoNotContact.value == "boss@ts.com"))


def test_a_file_cant_be_deleted_while_its_batch_runs(
    client: TestClient, owner: User, db: Session
) -> None:
    preview = upload(client, workbook(FILE))
    assert client.delete(f"/api/imports/{preview['id']}").json()["error"]["code"] == "waiting"
    client.post(f"/api/imports/{preview['id']}/confirm")
    batch = db.get(Batch, preview["batch_id"])
    assert batch is not None
    batch.status = "running"
    db.commit()

    refused = client.delete(f"/api/imports/{preview['id']}")

    assert (refused.status_code, refused.json()["error"]["code"]) == (409, "batch_running")
    assert client.get(f"/api/imports/{preview['id']}/deletion").json()["batch_running"] is True
    batch.status = "paused"
    db.commit()
    assert client.delete(f"/api/imports/{preview['id']}").status_code == 204


def test_viewers_cannot_delete_files_or_switch_cities(
    client: TestClient, owner: User, make_user: UserFactory
) -> None:
    preview = upload(client, workbook(FILE))
    client.post(f"/api/imports/{preview['id']}/confirm")
    city = client.get(f"/api/batches/{preview['batch_id']}/cities").json()[0]
    viewer = make_user(Role.VIEWER, name="Vic")
    sign_in(client, viewer.email)

    assert client.get(f"/api/batches/{preview['batch_id']}/cities").status_code == 200
    assert client.delete(f"/api/imports/{preview['id']}").status_code == 403
    assert client.patch(f"/api/lists/{city['id']}", json={"active": False}).status_code == 403
    assert (
        client.put(f"/api/batches/{preview['batch_id']}/cities", json={"active": False}).status_code
        == 403
    )


# --- small cities muted (step 16) -------------------------------------------------------------


def test_small_cities_are_muted_and_follow_new_files_unless_set_by_hand(
    client: TestClient, owner: User
) -> None:
    big = [row(f"Big {n}", f"https://share.google/b{n}", str(n)) for n in range(10)]
    first = upload(
        client,
        workbook(
            {
                "Tulsa": [STANDARD, *big],
                "Dallas": [
                    STANDARD,
                    *[
                        row(f"D {n}", f"https://share.google/d{n}", f"9{n}", city="Dallas")
                        for n in range(3)
                    ],
                ],
                "Austin": [STANDARD, row("A", "https://share.google/a", "71", city="Austin")],
            }
        ),
    )
    client.post(f"/api/imports/{first['id']}/confirm")
    url = f"/api/batches/{first['batch_id']}/cities"
    cities = {c["city"]: c for c in client.get(url).json()}
    assert {k: (c["active"], c["small"]) for k, c in cities.items()} == {
        "Austin": (False, True),
        "Dallas": (False, True),
        "Tulsa": (True, False),
    }
    # Unmuted by hand: the rule leaves Austin alone from now on.
    client.patch(f"/api/lists/{cities['Austin']['id']}", json={"active": True})

    more = [row(f"E {n}", f"https://share.google/e{n}", f"8{n}", city="Dallas") for n in range(7)]
    second = upload(client, workbook({"Dallas": [STANDARD, *more]}))
    client.post(f"/api/imports/{second['id']}/confirm")
    after = {c["city"]: c for c in client.get(url).json()}
    assert (after["Dallas"]["active"], after["Dallas"]["small"]) == (True, False)  # grew to 10
    assert after["Austin"]["active"] is True and after["Austin"]["set_by_hand"] is True

    # Deleting the second file makes Dallas small again: muted.
    client.delete(f"/api/imports/{second['id']}")
    assert {c["city"]: c["active"] for c in client.get(url).json()}["Dallas"] is False


def test_all_on_works_on_the_tab_shown(client: TestClient, owner: User) -> None:
    big = [row(f"Big {n}", f"https://share.google/b{n}", str(n)) for n in range(10)]
    preview = upload(
        client,
        workbook(
            {
                "Tulsa": [STANDARD, *big],
                "Austin": [STANDARD, row("A", "https://share.google/a", "71", city="Austin")],
            }
        ),
    )
    client.post(f"/api/imports/{preview['id']}/confirm")
    url = f"/api/batches/{preview['batch_id']}/cities"

    client.put(url, json={"active": False, "small": False})  # Cities (10+): all off
    client.put(url, json={"active": True, "small": True})  # Small cities: all on

    assert {c["city"]: c["active"] for c in client.get(url).json()} == {
        "Austin": True,
        "Tulsa": False,
    }
