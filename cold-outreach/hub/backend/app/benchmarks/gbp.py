"""Reading the GBP Benchmark file (PLAN.md section 3.5): one row per City + Niche, every
value optional. Every sheet of the workbook is read; the header row can be any of the first
five rows."""

import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.imports.reader import ReadError, clean

NUMBERS = (
    "additional_categories",
    "services",
    "photos",
    "posts_90_days",
    "description_length",
    "reviews",
    "rating",
    "review_reply_pct",
)

HEADERS: dict[str, tuple[str, ...]] = {
    "city": ("city",),
    "niche": ("niche",),
    "additional_categories": ("additional categories", "categories", "secondary categories"),
    "services": ("services", "service count", "number of services"),
    "photos": ("photos", "photo count", "number of photos", "images"),
    "posts_90_days": ("posts 90 days", "posts (90 days)", "posts in 90 days", "posts"),
    "description_length": ("description length", "description", "description characters"),
    "reviews": ("reviews", "number of reviews", "review count"),
    "rating": ("rating", "average rating", "average"),
    "review_reply_pct": ("review reply %", "review reply", "reply %", "reply rate", "replied %"),
    "reference_businesses": ("reference businesses", "businesses", "top businesses", "names"),
}


def _squash(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


HEADER_TO_FIELD = {_squash(name): f for f, names in HEADERS.items() for name in names}
LIMITS: dict[str, tuple[Decimal, Decimal]] = {
    "rating": (Decimal(0), Decimal(5)),
    "review_reply_pct": (Decimal(0), Decimal(100)),
}


@dataclass
class Report:
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ignored_columns: list[str] = field(default_factory=list)


def number(field_name: str, value: Any) -> Decimal | None:
    """A cell as a number. "85%" and 0.85 both mean 85 for the reply rate."""
    text = clean(value).replace(",", "")
    if not text:
        return None
    percent = text.endswith("%")
    try:
        result = Decimal(text.rstrip("%").strip())
    except InvalidOperation:
        raise ValueError("isn't a number") from None
    if field_name == "review_reply_pct" and not percent and result <= 1:
        result *= 100  # 0.85 written as a fraction
    low, high = LIMITS.get(field_name, (Decimal(0), Decimal(1_000_000)))
    if not low <= result <= high:
        raise ValueError(f"must be between {low} and {high}")
    return result.quantize(Decimal("0.01"))


def read(content: bytes) -> Report:
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        message = "This isn't an Excel file (.xlsx). Save it as .xlsx and try again."
        raise ReadError(message) from error
    report = Report()
    seen: dict[tuple[str, str], str] = {}
    for sheet in book.worksheets:
        cells = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        header_at = next(
            (
                i
                for i, row in enumerate(cells[:5])
                if {"city", "niche"} <= {HEADER_TO_FIELD.get(_squash(clean(c))) for c in row}
            ),
            None,
        )
        if header_at is None:
            continue
        columns: dict[int, str] = {}
        for index, cell in enumerate(cells[header_at]):
            name = clean(cell)
            target = HEADER_TO_FIELD.get(_squash(name)) if name else None
            if target and target not in columns.values():
                columns[index] = target
            elif name and name not in report.ignored_columns:
                report.ignored_columns.append(name)
        for line, raw in enumerate(cells[header_at + 1 :], start=header_at + 2):
            where = f"{clean(sheet.title)} row {line}"
            values = {t: (raw[i] if i < len(raw) else None) for i, t in columns.items()}
            if not any(clean(v) for v in values.values()):
                continue
            city, niche = clean(values.get("city")), clean(values.get("niche"))
            if not city or not niche:
                report.errors.append(f"{where}: City and Niche are both needed.")
                continue
            row: dict[str, Any] = {
                "city": city,
                "niche": niche,
                "reference_businesses": clean(values.get("reference_businesses"))[:1000],
            }
            bad = False
            for name in NUMBERS:
                try:
                    row[name] = number(name, values.get(name))
                except ValueError as problem:
                    report.errors.append(f"{where}: {name.replace('_', ' ')} {problem}.")
                    bad = True
            if bad:
                continue
            pair = (city.lower(), niche.lower())
            if pair in seen:
                report.errors.append(f"{where}: {city} · {niche} is also in {seen[pair]}.")
                continue
            seen[pair] = where
            report.rows.append(row)
    if not report.rows and not report.errors:
        raise ReadError("No sheet has a header row with City and Niche.")
    return report
