"""Fix files (SPEC step 28): the leads with a Lead Check problem as an Excel file, one tab
per city, each wrong cell red (an error) or yellow (a warning) with a note saying what is
wrong; fixed offline and uploaded back. Rows are matched by Lead ID, so a fix never lands on
the wrong business, and an empty cell clears the field (this file is the Hub's own)."""

import io
import re
import zipfile
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from app.db.models import Lead
from app.imports import check, reader
from app.leads.edit import FIELDS

LABELS = {**check.STANDARD, "notes": "Notes"}
HEADER = ["Lead ID", *(LABELS[f] for f in FIELDS), "Problems"]
RED = PatternFill("solid", fgColor="F8C9C9")
YELLOW = PatternFill("solid", fgColor="FFF1B8")
MAX_ROWS = 50_000


def tab_name(city: str, taken: set[str]) -> str:
    """A city as an Excel tab name: at most 31 characters, none of []:*?/\\, unique."""
    base = re.sub(r"[\[\]:*?/\\]", " ", city).strip()[:31] or "No city"
    name, n = base, 1
    while name.lower() in taken:
        n += 1
        name = f"{base[: 31 - len(str(n)) - 1]} {n}"
    taken.add(name.lower())
    return name


def build(leads: list[Lead]) -> bytes:
    """The fix file for these leads (each with its current values and problems)."""
    book = Workbook()
    first = book.active
    by_city: dict[str, list[Lead]] = {}
    for lead in leads[:MAX_ROWS]:
        by_city.setdefault(lead.city, []).append(lead)
    taken: set[str] = set()
    for city in sorted(by_city, key=str.lower):
        sheet = book.create_sheet(tab_name(city, taken))
        sheet.append(HEADER)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for lead in by_city[city]:
            values = check.lead_values(lead) | {"notes": (lead.data or {}).get("notes", "")}
            issues = lead.check or []
            sheet.append(
                [
                    lead.lead_id,
                    *(values.get(f, "") for f in FIELDS),
                    "; ".join(
                        f"{LABELS.get(i['field'], i['field'])}: {i['message']}" for i in issues
                    ),
                ]
            )
            line = sheet.max_row
            for cell in sheet[line]:
                cell.data_type = "s"  # text: a value starting with = never runs as a formula
            for item in issues:
                if item["field"] not in FIELDS:
                    continue
                cell = sheet.cell(line, FIELDS.index(item["field"]) + 2)
                error = item["severity"] == "error"
                if error or cell.fill != RED:
                    cell.fill = RED if error else YELLOW
                old = cell.comment.text + "\n" if cell.comment else ""
                cell.comment = Comment(
                    f"{old}{'Error' if error else 'Warning'}: {item['message']}", "Lead Check"
                )
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate([11, 28, 24, 34, 16, 28, 16, 10, 28, 28, 28, 9, 8, 24, 50]):
            sheet.column_dimensions[sheet.cell(1, index + 1).column_letter].width = width
    if first is not None and len(book.worksheets) > 1:
        book.remove(first)
    elif first is not None:
        first.title = "No problems"
        first.append(["Every lead here passes Lead Check."])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@dataclass
class FixRow:
    tab: str
    line: int
    lead_id: str
    values: dict[str, str] = field(default_factory=dict)  # only the file's columns


class FixFileError(ValueError):
    pass


def field_of(header: str) -> str | None:
    text = reader.clean(header).lower()
    for key, label in LABELS.items():
        if text == label.lower():
            return key
    found = reader.HEADER_TO_FIELD.get(re.sub(r"[^a-z0-9]", "", text))
    return found if found in FIELDS else None


def read(content: bytes) -> list[FixRow]:
    """Every row with a Lead ID, with the values of the file's columns ("" clears)."""
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise FixFileError("This isn't an Excel file (.xlsx).") from error
    rows: list[FixRow] = []
    found_id = False
    for sheet in book.worksheets:
        cells = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        if not cells:
            continue
        header = [reader.clean(c) for c in cells[0]]
        if "lead id" not in [h.lower() for h in header]:
            continue
        found_id = True
        id_at = [h.lower() for h in header].index("lead id")
        columns: dict[int, str] = {}
        for index, name in enumerate(header):
            target = field_of(name) if index != id_at else None
            if target and target not in columns.values():
                columns[index] = target
        for line, raw in enumerate(cells[1:], start=2):
            lead_id = reader.clean(raw[id_at] if id_at < len(raw) else "").upper()
            values = {
                target: reader.clean(raw[index] if index < len(raw) else "")
                for index, target in columns.items()
            }
            if not lead_id and not any(values.values()):
                continue
            rows.append(FixRow(sheet.title, line, lead_id, values))
            if len(rows) > MAX_ROWS:
                raise FixFileError(f"At most {MAX_ROWS:,} rows in one fix file.")
    if not found_id:
        raise FixFileError(
            "No Lead ID column: upload the fix file the Hub made (Download problems to fix)."
        )
    return rows
