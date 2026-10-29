"""Files handed in on the Files screen (SPEC step 31): Lead Check on a whole file before it
is in a batch, the file with each wrong cell marked, and a fingerprint of its content so the
same leads are never added to a batch twice; its cities and their Top businesses for Top
Review (step 32). Pure functions over the workbook."""

import hashlib
import io
import json
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.auth.tokens import now
from app.db.models import Submission
from app.imports import check, fixes, reader
from app.leads.edit import FIELDS as EDIT_FIELDS

REPORT_TAB = "Lead Check"  # the tab the marked file opens on; taken out again on upload
REPORT_HEADER = ["Severity", "Tab", "Row", "Column", "Business", "Problem", "Value"]
AUTHOR = "Lead Check"
RED = PatternFill("solid", fgColor="F8C9C9")
YELLOW = PatternFill("solid", fgColor="FFF1B8")
NO_FILL = PatternFill(fill_type=None)


@dataclass
class Checked:
    businesses: int
    fingerprint: str
    # errors, warnings, each tab's counts and column problems, and the cities (step 32)
    check: dict[str, Any]
    problems: list[dict[str, Any]] = field(default_factory=list)
    # Each business for the Files table (step 33): its tab and row, values and problems.
    leads: list[dict[str, Any]] = field(default_factory=list)


def _fingerprint(report: reader.Report) -> str:
    """The same leads give the same fingerprint, however the file was saved."""
    rows = [[row.tab, sorted(row.values.items())] for row in report.sheet]
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode()).hexdigest()


def _where(lead: reader.Lead, field_name: str, rows: dict[str, reader.SheetRow]) -> reader.SheetRow:
    """The row of a merged business whose cell holds the value with the problem."""
    found = [rows[w] for w in lead.rows if w in rows]
    for row in found:
        if row.values.get(field_name):
            return row
    return found[0]


TABLE_FIELDS = (*check.STANDARD, "notes")


@dataclass
class Others:
    """What the file is checked against (step 39), read from the database before: the
    cities another file of its niche (handed in earlier) or its batch already has, each with
    why this file's is muted, and the businesses already handed in or in a batch."""

    cities: dict[str, str] = field(default_factory=dict)  # city (lower case) -> why muted
    # The values of a business -> where the same business already is, or None.
    same: Callable[[dict[str, str]], str | None] = lambda _values: None


def run(
    content: bytes,
    niche: str,
    progress: Callable[[int], None] | None = None,
    number: int = 0,
    others: Others | None = None,
) -> Checked:
    """Lead Check on a file: every problem with its tab, row, column and value. A business
    given on several rows is checked once, as it will be saved, and is an error (step 39), as
    is a business already handed in (`others`). A muted city's businesses aren't checked.
    `progress` is told how far it is, in percent: reading the tabs is the first 60, checking
    the businesses the rest. Each business gets its File ID (step 35): the one in the file if
    the Hub made it, else the next of the file `number`."""

    def tell(percent: float) -> None:
        if progress is not None:
            progress(min(99, int(percent)))

    others = others or Others()
    report = reader.read(content, niche, lambda share: tell(share * 60))
    found_cities = _cities(report, others.cities)
    muted = {key for key, item in found_cities.items() if item["muted"]}
    rows = {f"{row.tab} row {row.line}": row for row in report.sheet}
    # A tab whose businesses are all in muted cities isn't checked either.
    live_tabs = {
        rows[lead.rows[0]].tab
        for lead in report.leads
        if lead.rows[0] in rows and city_key(lead.values.get("city", "")) not in muted
    }
    tabs: dict[str, dict[str, Any]] = {}
    problems: list[dict[str, Any]] = []
    for name in report.tab_names:
        layout = report.tab_columns.get(name)
        columns = (
            check.check_tab(layout["fields"], layout["duplicates"])
            if layout and (name in live_tabs or not muted)
            else []
        )
        tabs[name] = {"errors": 0, "warnings": 0, "columns": [i.as_dict() for i in columns]}
    order = {name: index for index, name in enumerate(report.tab_names)}
    leads: list[dict[str, Any]] = []
    for index, lead in enumerate(report.leads):
        if index % 200 == 0:
            tell(60 + 40 * index / max(1, len(report.leads)))
        is_muted = city_key(lead.values.get("city", "")) in muted
        found = [] if is_muted else check.check_lead(lead.values) + _same(lead, others)
        first = rows.get(lead.rows[0])
        leads.append(
            {
                "tab": first.tab if first else "",
                "line": first.line if first else 0,
                "rows": lead.rows,
                "values": {k: lead.values[k][:500] for k in TABLE_FIELDS if lead.values.get(k)},
                "extra": lead.extra,
                "issues": [i.as_dict() for i in found],
                "id": lead.values.get("file_id", ""),
            }
            | ({"muted": True} if is_muted else {})
        )
        for item in found:
            row = _where(lead, item.field, rows)
            kind = "errors" if item.severity == "error" else "warnings"
            tabs.setdefault(row.tab, {"errors": 0, "warnings": 0, "columns": []})[kind] += 1
            problems.append(
                item.as_dict()
                | {
                    "tab": row.tab,
                    "line": row.line,
                    "business": lead.values.get("business_name") or lead.values.get("gbp_link"),
                    "value": lead.values.get(item.field, "")[:300],
                }
            )
    _give_ids(leads, number)
    problems.sort(key=lambda p: (order.get(p["tab"], 0), p["line"], p["severity"] != "error"))
    # A tab's column problems count by their severity (errors since step 36).
    errors = sum(
        t["errors"] + sum(1 for c in t["columns"] if c["severity"] == "error")
        for t in tabs.values()
    )
    warnings = sum(
        t["warnings"] + sum(1 for c in t["columns"] if c["severity"] != "error")
        for t in tabs.values()
    )
    return Checked(
        businesses=len(report.leads),
        fingerprint=_fingerprint(report),
        check={"errors": errors, "warnings": warnings, "tabs": tabs, "cities": found_cities},
        problems=problems,
        leads=leads,
    )


def _same(lead: reader.Lead, others: Others) -> list[check.Issue]:
    """The same business twice (step 39): on several rows of the file, or handed in before."""
    found = []
    if len(lead.rows) > 1:
        severity, message = check.CODES["SAME_IN_FILE"]
        where = ", ".join(lead.rows)
        found.append(
            check.Issue(
                "SAME_IN_FILE",
                "business_name",
                severity,
                f"{message} ({where}): press Remove duplicates",
            )
        )
    elsewhere = others.same(lead.values)
    if elsewhere:
        severity, message = check.CODES["SAME_ELSEWHERE"]
        found.append(
            check.Issue("SAME_ELSEWHERE", "business_name", severity, f"{message}: {elsewhere}")
        )
    return found


# --- Top Review (step 32) -------------------------------------------------------------------


def city_key(city: str) -> str:
    return city.strip().lower()


def _cities(report: reader.Report, muted: dict[str, str]) -> dict[str, dict[str, Any]]:
    """Each city of the file (lower case) -> its name, how many businesses it has, and why
    it is muted (empty when it isn't)."""
    found: dict[str, dict[str, Any]] = {}
    for lead in report.leads:
        city = lead.values.get("city", "")
        key = city_key(city)
        item = found.setdefault(key, {"name": city, "businesses": 0, "muted": muted.get(key, "")})
        item["businesses"] += 1
    return found


FINAL_CHECK = timedelta(hours=1)  # a file without errors waits this long before Top Review
READY_WAIT = timedelta(hours=1)  # a city with a Top waits this long before its batch


def seconds_until(when: datetime | None) -> int | None:
    """How long until then, by the server's clock (never the viewer's, which may be off)."""
    return None if when is None else max(0, int((when - now()).total_seconds()))


def send_on(sub: Submission) -> None:
    """The Final check is over (step 36): the file goes on to Top Review."""
    sub.sent_at = now()


def city_businesses(content: bytes, niche: str, key: str) -> list[dict[str, Any]]:
    """The businesses of one city of the file, as they will be saved: each with its key
    (the first row it is on), its values and its other columns."""
    report = reader.read(content, niche)
    return [
        {
            "key": lead.rows[0],
            "rows": lead.rows,
            "values": {k: v for k, v in lead.values.items() if v and k not in ("niche", "level")},
            "extra": lead.extra,
        }
        for lead in report.leads
        if city_key(lead.values.get("city", "")) == key
    ]


def _open(content: bytes) -> Any:
    try:
        return load_workbook(io.BytesIO(content))
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        message = "This isn't an Excel file (.xlsx). Save it as .xlsx and try again."
        raise reader.ReadError(message) from error


def _unmark(book: Any) -> bool:
    """Take out what an earlier marked copy added: its report tab, notes and colours."""
    changed = False
    for sheet in list(book.worksheets):
        if sheet.title.startswith(REPORT_TAB) and sheet.cell(1, 1).value == REPORT_HEADER[0]:
            book.remove(sheet)
            changed = True
            continue
        for line in sheet.iter_rows():
            for cell in line:
                if cell.comment is not None and cell.comment.author == AUTHOR:
                    cell.comment = None
                    if cell.fill.fgColor.rgb in ("00F8C9C9", "00FFF1B8"):
                        cell.fill = NO_FILL
                    changed = True
    return changed


def clean(content: bytes) -> bytes:
    """A revised file as handed in, without the marks of the copy it was fixed in."""
    book = _open(content)
    if not _unmark(book):
        return content
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --- the fix file (steps 34 and 35) -----------------------------------------------------------
# Like a batch's fix file (step 28), on one sheet: the businesses with a problem by city, each
# found again by its File ID, every standard column with each wrong cell marked, the problems,
# and Delete = Yes to take it out of the file.

FIX_FIELDS = EDIT_FIELDS  # Business name … Notes, as the batch's fix file has them
FILE_ID, PROBLEMS, DELETE = "File ID", "Problems", "Delete"
FIX_HEADER = [FILE_ID, *(fixes.LABELS[f] for f in FIX_FIELDS), PROBLEMS, DELETE]
FIX_SHEET = "Problems"
# The same business twice (step 39): never in the fix file, removed on the Hub (step 40).
DUPLICATE_CODES = ("SAME_IN_FILE", "SAME_ELSEWHERE")
YES_WORDS = ("yes", "y", "x", "delete")
FILE_ID_FORM = re.compile(r"^F(\d+)-(\d{4,})$")


def file_id(number: int, n: int) -> str:
    return f"F{number}-{n:04d}"


def _give_ids(leads: list[dict[str, Any]], number: int) -> None:
    """Businesses keep a File ID of this file the Hub wrote in it (once each); the others
    get the next ones, in the file's order."""
    taken: set[str] = set()
    highest = 0
    for lead in leads:
        found = FILE_ID_FORM.match(lead["id"].upper())
        if found and int(found[1]) == number and lead["id"].upper() not in taken:
            lead["id"] = lead["id"].upper()
            taken.add(lead["id"])
            highest = max(highest, int(found[2]))
        else:
            lead["id"] = ""
    for lead in leads:
        if not lead["id"]:
            highest += 1
            lead["id"] = file_id(number, highest)


def fix_file(leads: list[dict[str, Any]]) -> bytes:
    """The businesses with a problem, on one sheet by city, to fix in Excel and upload."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = FIX_SHEET
    sheet.append(FIX_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    # Duplicates aren't in it: they are removed on the Hub only (step 40).
    to_fix = [
        (lead, [i for i in lead["issues"] if i["code"] not in DUPLICATE_CODES]) for lead in leads
    ]
    shown = sorted(
        ((lead, issues) for lead, issues in to_fix if issues),
        key=lambda item: (item[0]["values"].get("city", "").lower(), item[0]["id"]),
    )
    for lead, issues in shown:
        sheet.append(
            [
                lead["id"],
                *(lead["values"].get(f, "") for f in FIX_FIELDS),
                "; ".join(
                    f"{fixes.LABELS.get(i['field'], i['field'])}: {i['message']}" for i in issues
                ),
                "",
            ]
        )
        line = sheet.max_row
        for cell in sheet[line]:
            cell.data_type = "s"  # text: a value starting with = never runs as a formula
        for item in issues:
            if item["field"] not in FIX_FIELDS:
                continue
            cell = sheet.cell(line, FIX_FIELDS.index(item["field"]) + 2)
            error = item["severity"] == "error"
            if error or cell.fill != RED:
                cell.fill = RED if error else YELLOW
            old = cell.comment.text + "\n" if cell.comment else ""
            cell.comment = Comment(
                f"{old}{'Error' if error else 'Warning'}: {item['message']}", AUTHOR
            )
    if not shown:
        duplicates = any(len(issues) < len(lead["issues"]) for lead, issues in to_fix)
        sheet.append(
            [
                "",
                "Nothing to fix here: remove the duplicates on the Hub (Remove duplicates)."
                if duplicates
                else "Every business in this file passes Lead Check.",
            ]
        )
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [11, 28, 24, 34, 16, 28, 16, 10, 28, 28, 28, 9, 8, 24, 50, 8]
    for index, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(index + 1)].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class NoFileId(fixes.FixFileError):
    """The file has no File ID column: a whole new version of the file, not a fix file."""


@dataclass
class FixRow:
    tab: str  # where in the fix file, for messages
    line: int
    key: str  # its File ID: "F12-0014"
    values: dict[str, str]  # only the fix file's columns ("" clears)
    delete: bool


def _cell(raw: tuple[Any, ...], at: int | None) -> str:
    return reader.clean(raw[at]) if at is not None and at < len(raw) else ""


def read_fixes(content: bytes) -> list[FixRow]:
    """Every row of a fix file with a File ID, with its columns' values and Delete."""
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise fixes.FixFileError("This isn't an Excel file (.xlsx).") from error
    rows: list[FixRow] = []
    found = False
    for sheet in book.worksheets:
        cells = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        if not cells:
            continue
        header = [reader.clean(c).lower() for c in cells[0]]
        if FILE_ID.lower() not in header:
            continue
        found = True
        key_at = header.index(FILE_ID.lower())
        delete_at = header.index(DELETE.lower()) if DELETE.lower() in header else None
        columns: dict[int, str] = {}
        for index, name in enumerate(header):
            target = fixes.field_of(name) if index not in (key_at, delete_at) else None
            if target and target not in columns.values():
                columns[index] = target
        for line, raw in enumerate(cells[1:], start=2):
            values = {target: _cell(raw, index) for index, target in columns.items()}
            key = _cell(raw, key_at).upper()
            if not key and not any(values.values()):
                continue
            delete = _cell(raw, delete_at).lower() in YES_WORDS
            rows.append(FixRow(sheet.title, line, key, values, delete))
            if len(rows) > fixes.MAX_ROWS:
                raise fixes.FixFileError(f"At most {fixes.MAX_ROWS:,} rows in one fix file.")
    if not found:
        raise NoFileId("No File ID column: this is a whole file, not the Hub's fix file.")
    return rows


def rebuilt(leads: list[dict[str, Any]], once: bool = False) -> bytes:
    """The file made again from its businesses (after fixes): a tab per tab of the file, each
    business's File ID (kept when it is checked again), the standard columns and the file's
    other columns. A business on several rows is written on each of them, so it stays a
    duplicate, unless `once` (Remove duplicates, step 40): then on its first row only."""
    book = Workbook()
    first = book.active
    extra: list[str] = []
    for lead in leads:
        extra += [name for name in lead.get("extra", {}) if name not in extra]
    header = [FILE_ID, *(fixes.LABELS[f] for f in FIX_FIELDS), *extra]
    sheets: dict[str, Any] = {}
    taken: set[str] = set()
    placed = [
        (lead, where.rsplit(" row ", 1)[0] if " row " in where else lead["tab"])
        for lead in leads
        for where in (lead.get("rows") or [""])[: 1 if once else None]
    ]
    for lead, tab_of in placed:
        tab = tab_of or lead["tab"] or lead["values"].get("city", "") or "Leads"
        if tab not in sheets:
            sheets[tab] = book.create_sheet(fixes.tab_name(tab, taken))
            sheets[tab].append(header)
        values = lead["values"]
        sheets[tab].append(
            [
                lead.get("id", ""),
                *(values.get(f, "") for f in FIX_FIELDS),
                *(lead.get("extra", {}).get(name, "") for name in extra),
            ]
        )
        for cell in sheets[tab][sheets[tab].max_row]:
            cell.data_type = "s"
    if first is not None and sheets:
        book.remove(first)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
