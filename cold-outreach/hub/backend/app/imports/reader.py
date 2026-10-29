"""Reading a lead file (SPEC section 6.4): one tab per city, the header spellings the team
already uses, cleaning, and merging duplicate rows into one lead. Pure functions over the
workbook, so every rule is easy to test."""

import io
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any
from urllib.parse import parse_qsl, unquote, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

# Lead fields, in the order the Leads tab shows them (PLAN.md section 3.2).
FIELDS = (
    "city",
    "niche",
    "level",
    "business_name",
    "gbp_link",
    "address",
    "website",
    "phone",
    "whatsapp",
    "facebook",
    "instagram",
    "email",
    "reviews",
    "rating",
    "notes",
)

# Header spellings seen in the team's files, compared without case, spaces or punctuation.
SYNONYMS: dict[str, tuple[str, ...]] = {
    "business_name": ("business name", "business", "name", "company", "company name"),
    "gbp_link": ("map", "map url", "map link", "maps", "google map", "gbp", "gbp link"),
    "address": ("address", "full address"),
    "city": ("city", "town"),
    "website": ("website", "web site", "site", "url"),
    "phone": ("phone", "phone number", "telephone", "tel"),
    "whatsapp": ("wa", "wp", "whatsapp", "whats app"),
    "facebook": ("facebook", "fb"),
    "instagram": ("instagram", "ig", "insta"),
    "email": ("email", "e-mail", "emails", "gmail", "mail"),
    "reviews": ("reviews", "review", "number of reviews", "reviews count", "review count"),
    "rating": ("rating", "ratings", "average", "average rating", "average ratings"),
    "notes": ("notes", "note", "comment", "comments", "issue 1", "issue 2", "issue 3"),
}
# Top is chosen in the Hub (SPEC section 6.4 item 5): a Level column is kept like any other.


def _squash(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


HEADER_TO_FIELD = {_squash(name): f for f, names in SYNONYMS.items() for name in names}
# The File ID the Hub gives each business of a handed-in file (step 35), in the files it
# makes: read back so a business keeps it, never a lead field of its own.
HEADER_TO_FIELD["fileid"] = "file_id"

EMAIL = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[a-z]{2,}$", re.IGNORECASE)
NOT_OWN_SITE = re.compile(r"(facebook|instagram|linktr\.ee|yelp)\.", re.IGNORECASE)


class ReadError(Exception):
    """The file can't be read at all (not an Excel file, or no tabs with a header)."""


@dataclass
class Row:
    tab: str
    line: int  # the row number in Excel, for messages
    values: dict[str, str]
    emails: list[str]
    extra: dict[str, str] = field(default_factory=dict)  # other columns, by their header


@dataclass
class SheetRow:
    """A row of the file as the preview shows it, tab by tab (SPEC section 6.4)."""

    tab: str
    line: int
    values: dict[str, str]  # cleaned, by field (emails joined)
    extra: dict[str, str]
    raw: dict[str, str]  # the original text of cells the Hub changed, by field
    bad_emails: list[str]
    skipped: bool = False  # no name and no GBP link


@dataclass
class Lead:
    """One business after merging. `rows` says where it came from."""

    values: dict[str, str]
    emails: list[str]
    rows: list[str] = field(default_factory=list)  # "Tulsa row 14"
    extra: dict[str, str] = field(default_factory=dict)


@dataclass
class Report:
    leads: list[Lead]
    tabs: int = 0
    empty_tabs: list[str] = field(default_factory=list)
    unknown_columns: dict[str, list[str]] = field(default_factory=dict)  # column -> tabs
    skipped_rows: list[str] = field(default_factory=list)  # no name and no GBP link
    invalid_emails: list[str] = field(default_factory=list)  # "Tulsa row 3: kevin@"
    merged_rows: list[list[str]] = field(default_factory=list)  # each group of merged rows
    has_level_column: bool = False
    sheet: list[SheetRow] = field(default_factory=list)  # every row with a business, by tab
    tab_names: list[str] = field(default_factory=list)  # every tab, in the file's order
    extra_columns: list[str] = field(default_factory=list)  # other headers, first seen first
    # Each tab's standard columns and the ones given twice, for Lead Check (step 26):
    # tab -> {"fields", "duplicates"}.
    tab_columns: dict[str, dict[str, Any]] = field(default_factory=dict)


# --- cleaning -------------------------------------------------------------------------------


def clean(value: Any) -> str:
    """Cell text without surrounding spaces or line breaks; numbers without a trailing .0."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def phone_digits(phone: str) -> str:
    """The last 10 digits, so "(918) 523-8300" and "+1 918 523 8300" match."""
    digits = re.sub(r"\D", "", phone)
    return digits[-10:] if len(digits) >= 10 else ""


def clean_website(url: str) -> str:
    """Remove utm_* tracking parameters; keep everything else as given."""
    if not url or "utm_" not in url:
        return url
    parts = urlsplit(url)
    query = [
        (k, v)
        for k, v in parse_qsl(parts.query, keep_blank_values=True)
        if not k.startswith("utm_")
    ]
    return urlunsplit(parts._replace(query=urlencode(query)))


def clean_gbp(url: str) -> str:
    return url.strip().rstrip("/")


def website_domain(url: str) -> str:
    host = urlsplit(url if "//" in url else f"//{url}").netloc.lower()
    return host.removeprefix("www.")


def split_emails(text: str) -> tuple[list[str], list[str]]:
    """(valid, invalid) emails from a cell that may hold several."""
    valid: list[str] = []
    invalid: list[str] = []
    for part in re.split(r"[\s,;/]+", unquote(text)):
        email = part.strip().strip(".").lower()
        if not email:
            continue
        (valid if EMAIL.match(email) else invalid).append(email)
    return valid, invalid


def normal_text(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def name_place(values: dict[str, str]) -> str:
    """Name + address; empty without either (step 39: a name alone, or a name and a city,
    is never enough, as one brand has a business in many cities)."""
    name = normal_text(values.get("business_name", ""))
    address = normal_text(values.get("address", ""))
    return f"{name}|{address}" if name and address else ""


def keys(values: dict[str, str]) -> set[str]:
    """What makes two rows the same business: the GBP link, or name + address."""
    found = set()
    if values.get("gbp_link"):
        found.add("gbp:" + values["gbp_link"])
    place = name_place(values)
    if place:
        found.add("name:" + place)
    return found


def is_top(level: str) -> bool:
    return level.strip().lower() == "top"


# --- reading --------------------------------------------------------------------------------


def _header_row(rows: list[tuple[Any, ...]]) -> int | None:
    """The first row (of the first five) that has at least two known headers."""
    for index, row in enumerate(rows[:5]):
        known = {HEADER_TO_FIELD.get(_squash(clean(cell))) for cell in row if clean(cell)}
        known.discard(None)
        if len(known) >= 2:
            return index
    return None


def read_rows(
    content: bytes, report: Report, progress: Callable[[float], None] | None = None
) -> list[Row]:
    """`progress` is told the share of tabs read (0 to 1) after each tab."""
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        message = "This isn't an Excel file (.xlsx). Save it as .xlsx and try again."
        raise ReadError(message) from error
    rows: list[Row] = []
    total = len(book.worksheets)
    for done, sheet in enumerate(book.worksheets):
        if progress is not None:
            progress(done / total)
        report.tabs += 1
        tab = clean(sheet.title)
        report.tab_names.append(tab)
        cells = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        header_at = _header_row(cells)
        if header_at is None:
            report.empty_tabs.append(tab)
            continue
        columns: dict[int, str] = {}
        others: dict[int, str] = {}
        duplicates: list[str] = []
        for index, cell in enumerate(cells[header_at]):
            name = clean(cell)
            if not name:
                continue
            target = HEADER_TO_FIELD.get(_squash(name))
            if _squash(name) == "level":
                report.has_level_column = True
            if target is None:
                report.unknown_columns.setdefault(name, []).append(tab)
                if name not in others.values():
                    others[index] = name[:80]
                if name[:80] not in report.extra_columns:
                    report.extra_columns.append(name[:80])
            elif target not in columns.values() or target in ("notes", "email"):
                columns[index] = target
            elif target not in duplicates:
                duplicates.append(target)
        layout: dict[str, Any] = {
            "fields": sorted(set(columns.values())),
            "duplicates": duplicates,
        }
        report.tab_columns[tab] = layout
        found = 0
        for offset, raw in enumerate(cells[header_at + 1 :], start=header_at + 2):
            values: dict[str, str] = {}
            emails: list[str] = []
            bad: list[str] = []
            original: dict[str, str] = {}
            for index, target in columns.items():
                cell = raw[index] if index < len(raw) else None
                text = clean(cell)
                if not text:
                    continue
                if target == "email":
                    valid, invalid = split_emails(text)
                    emails.extend(e for e in valid if e not in emails)
                    bad.extend(invalid)
                    report.invalid_emails.extend(f"{tab} row {offset}: {x}" for x in invalid)
                    original["email"] = (original.get("email", "") + " " + str(cell)).strip()
                elif target == "notes" and values.get("notes"):
                    values["notes"] += " · " + text
                else:
                    values[target] = text
                    original[target] = str(cell).strip()
            extra = {
                name: text
                for index, name in others.items()
                if index < len(raw) and (text := clean(raw[index]))
            }
            # A city filled down an otherwise empty sheet isn't a business.
            if not emails and not bad and not set(values) - {"city", "level"}:
                continue
            found += 1
            skipped = not values.get("business_name") and not values.get("gbp_link")
            if not skipped:
                values["city"] = values.get("city") or tab
                values["website"] = clean_website(values.get("website", ""))
                values["gbp_link"] = clean_gbp(values.get("gbp_link", ""))
            shown = {k: v for k, v in values.items() if v}
            if emails:
                shown["email"] = ", ".join(emails)
            changed = {k: v[:500] for k, v in original.items() if v and v != shown.get(k, "")}
            report.sheet.append(SheetRow(tab, offset, shown, extra, changed, bad, skipped=skipped))
            if skipped:
                report.skipped_rows.append(f"{tab} row {offset}")
                continue
            rows.append(Row(tab=tab, line=offset, values=values, emails=emails, extra=extra))
        if not found:
            report.empty_tabs.append(tab)
    return rows


def same_business(a: dict[str, str], b: dict[str, str]) -> bool:
    """Two rows are one business when their GBP links match, or their business name and
    address both match (step 39). The same name in another city, or the same phone, isn't
    enough: one brand often has a business in many cities, each at its own address."""
    return bool(keys(a) & keys(b))


def merge(rows: list[Row], niche: str, report: Report) -> list[Lead]:
    """Rows of the same business become one lead: emails are pooled, gaps filled from the
    later rows, and the first row's value kept when both have one. Top wins."""
    leads: list[Lead] = []
    by_key: dict[str, list[Lead]] = {}
    for row in rows:
        values = dict(row.values, niche=niche)
        where = f"{row.tab} row {row.line}"
        candidates = [lead for k in keys(values) for lead in by_key.get(k, [])]
        match = next((lead for lead in candidates if same_business(lead.values, values)), None)
        if match is None:
            lead = Lead(values=values, emails=list(row.emails), rows=[where], extra=dict(row.extra))
            leads.append(lead)
        else:
            for name, value in row.extra.items():
                lead.extra.setdefault(name, value)
            lead = match
            for name, value in values.items():
                if value and not lead.values.get(name):
                    lead.values[name] = value
            if is_top(values.get("level", "")):
                lead.values["level"] = "Top"
            lead.emails.extend(e for e in row.emails if e not in lead.emails)
            lead.rows.append(where)
        for key in keys(lead.values):
            bucket = by_key.setdefault(key, [])
            if lead not in bucket:
                bucket.append(lead)
    for lead in leads:
        lead.values["level"] = "Top" if is_top(lead.values.get("level", "")) else ""
        lead.values["email"] = ", ".join(lead.emails)
        if len(lead.rows) > 1:
            report.merged_rows.append(lead.rows)
    return leads


def read(content: bytes, niche: str, progress: Callable[[float], None] | None = None) -> Report:
    report = Report(leads=[])
    rows = read_rows(content, report, progress)
    if report.tabs == len(report.empty_tabs):
        raise ReadError("No tab has a header row with columns such as Business name and MAP.")
    report.leads = merge(rows, niche, report)
    return report


def not_own_site(url: str) -> bool:
    return bool(NOT_OWN_SITE.search(url))
